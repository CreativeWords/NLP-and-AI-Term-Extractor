r"""
Terminology Extractor (Mixed Sources: TMX + PDF + DOCX)
-------------------------------------------------------
Builds bilingual glossaries from a project's source material. Extracts
source-language terminology candidates (any spaCy-supported source language,
including Chinese and Japanese) from a folder, optionally scanned
recursively, containing any mix of:
  - TMX translation memories  -> one segment per source-language TU, with
                                 its aligned target segment when available
  - PDF files                 -> sentences (image-only pages skipped)
  - DOCX files                -> sentences (from paragraphs + table cells)

Legacy .doc files are detected but skipped (python-docx cannot read them);
a count is logged so the user knows they were present.

Pipeline:
  1. Collect supported files; pre-check TMX source languages and resolve
     the TMX target language(s) to use
  2. Parse files into tagged segments (source text + optional aligned target)
  3. Clean & filter segments
  4. Extract candidates per segment using spaCy (noun chunks, POS-based
     term typing, example capture, optional junk filter)
  5. Aggregate by segment count, raw frequency, and source provenance;
     apply n-gram, min-count and term-cap filters (Chinese/Japanese use
     character length instead of word n-grams)
  6. (Optional) AI cleaning: whole-list merge/rename/drop via decisions-by-id
  7. (Optional) AI translation + definitions, batched, in a single step:
     each request returns the target term and a definition. Terms with
     TMX-aligned examples use an extraction prompt that reads the verified
     human translation
  8. Export to JSON + a mirrored flat XLSX (same schema)

Any AI step failure halts the AI chain; output reflects the last fully
completed stage.

A Tkinter GUI lets the user:
  - browse for a folder
  - choose the source language, n-gram lengths, term cap mode and advanced
    extraction options (junk filter, min. segment length, min. count,
    custom stopwords)
  - optionally enable AI cleaning and translation + definitions, pick the
    model (Anthropic or OpenAI), and supply an API key, domain, context,
    target language and definition language
  - run the pipeline, watching progress and per-step token usage in a log box

Requirements:
    pip install lxml spacy openpyxl pdfplumber python-docx
    # extra, only needed when AI enhancement is used:
    pip install anthropic openai pyyaml

The AI prompts live in an external file, 'term_extractor_prompts.yaml',
which must sit next to this script.

spaCy models are not installed manually: the tool downloads the large (lg)
model for the selected language on first use into a per-user cache folder
(see SPACY_MODELS_DIR; override with the TERM_EXTRACTOR_MODELS_DIR
environment variable) and loads it from there on subsequent runs.
"""

import re
import queue
import threading
import xml.etree.ElementTree as ET
from collections import defaultdict, Counter
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import tkinter as tk
from tkinter import filedialog, ttk, scrolledtext
from tkinter import font as tkfont

import spacy
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

import os
import sys
import json
import html


# ---------------------------------------------------------------------------
# CONFIGURATION
# ---------------------------------------------------------------------------

VERSION = "0.7.8"

# Provenance tags used in the 'sources' column for TMX segments whose target
# language could not be used (no match, or the user chose not to use it).
# The "AI-translated" variant is used only when the translation step will run,
# so the label stays accurate when it won't.
TAG_TMX_SOURCE_ONLY = "tmx (source only)"
TAG_TMX_SOURCE_ONLY_AI = "tmx (source only, AI-translated)"

# --- Junk candidate filtering -------------------------------------------
# Symbol units only: these are SI/technical notation and identical across
# languages, so the filter stays language-independent. Spelled-out durations
# ("2 years") are deliberately NOT listed — they differ per language and the
# AI cleaning step removes them anyway.
TERM_JUNK_UNITS = {
    "cm", "mm", "m", "km", "ml", "l", "kg", "g", "mg", "µg", "µm", "um", "nm",
    "º", "°", "%", "mmhg", "psi", "fr", "ch", "hz", "khz", "mhz", "v", "w",
    "a", "ma", "kpa", "bar", "in", "ft",
    # Cyrillic unit symbols (Russian / Ukrainian / Macedonian notation), so
    # glued measurements like "10мл шприц" are not mistaken for codes.
    # "в" (volt) is deliberately omitted: it is also a common preposition.
    "мм", "см", "м", "км", "мл", "л", "кг", "г", "мг", "мкг", "мкм", "нм",
    "вт", "гц", "кгц", "мгц", "кпа", "мпа", "па", "ма",
}

# --- Script classes for the junk filter ------------------------------------
# Python's \w is Unicode-aware, so [^\W\d_] is "any letter in any script".
# CJK scripts (Han, kana, Hangul) need different treatment from alphabetic
# ones: one or two characters are a normal term length ("弁", "安全"), and
# characters sitting next to digits are ordinary prose ("3号機", "100V電源"),
# not product codes. Codes in CJK documents are written in Latin letters and
# are still caught by the alphabetic rules.
_CJK_RANGES = (
    "\u1100-\u11ff"   # Hangul Jamo
    "\u3040-\u309f"   # Hiragana
    "\u30a0-\u30ff"   # Katakana
    "\u3130-\u318f"   # Hangul Compatibility Jamo
    "\u31f0-\u31ff"   # Katakana Phonetic Extensions
    "\u3400-\u4dbf"   # CJK Unified Ideographs Extension A
    "\u4e00-\u9fff"   # CJK Unified Ideographs
    "\uac00-\ud7af"   # Hangul Syllables
    "\uf900-\ufaff"   # CJK Compatibility Ideographs
    "\uff66-\uff9f"   # Halfwidth Katakana
)
_CJK_RE = re.compile(f"[{_CJK_RANGES}]")
_ANY_LETTER = r"[^\W\d_]"
# A letter of an alphabetic script (Latin, Cyrillic, Greek, ...): any letter
# that is not CJK. Used by the length and code-pattern rules.
_ALPHA_LETTER = rf"(?:(?![{_CJK_RANGES}]){_ANY_LETTER})"
_ANY_LETTER_RE = re.compile(_ANY_LETTER)
_ALPHA_PAIR_RE = re.compile(_ALPHA_LETTER + "{2}")
_CODE_PATTERN_RE = re.compile(
    rf"({_ALPHA_LETTER}{{2,}})-?(\d{{2,}})|(\d{{2,}})-?({_ALPHA_LETTER}{{2,}})")
# Characters that never legitimately open a term; a candidate starting with
# one is an extraction artifact (e.g. "[A", "(inches", "| LBL-02603").
TERM_JUNK_OPENERS = '([{|"\'«»<“”„‘’'

# Maximum examples to store per term
MAX_EXAMPLES = 3
# Larger pool of candidate examples kept per term before final selection, so
# selection can prefer TMX-aligned examples and merges can re-select over the
# combined pool. First-seen order is used, so a modest pool is enough.
EXAMPLE_POOL_MAX = 6

# --- AI (Anthropic / OpenAI) configuration ---
# Selectable models shown in the GUI: label -> (api_model_string, provider).
AI_MODELS = {
    "Claude Opus 5":   ("claude-opus-5", "anthropic"),
    "Claude Sonnet 5": ("claude-sonnet-5", "anthropic"),
    "GPT 5.6 Sol":     ("gpt-5.6-sol", "openai"),
    "GPT 5.6 Terra":   ("gpt-5.6-terra", "openai"),
    "GPT 5.6 Luna":    ("gpt-5.6-luna", "openai"),
}
DEFAULT_AI_MODEL_LABEL = "Claude Sonnet 5"
AI_MAX_TOKENS = 96000                 # output ceiling for JSON responses
# Transport-level retries performed by the provider SDKs (connection errors,
# 429 rate limits, 5xx) with their own backoff. Above the library defaults.
AI_SDK_MAX_RETRIES = 5
TRANSLATE_BATCH_SIZE = 40             # terms per translation + definition request


class AIStepFailure(Exception):
    """
    Raised when an AI step cannot complete (API error, unparseable JSON,
    or a response missing its expected array). The pipeline halts the AI
    chain on this and writes output from the last fully completed step —
    the failed step's partial work is discarded.

    Truncation is NOT a failure: it has its own salvage-and-continue
    behaviour and never raises this.
    """


def _model_provider(model: str) -> str:
    """Return 'anthropic' or 'openai' for an API model string."""
    for _label, (mstr, provider) in AI_MODELS.items():
        if mstr == model:
            return provider
    # Fallback heuristic if an unknown string is passed.
    return "openai" if model.startswith("gpt") else "anthropic"

# Per-model pricing, USD per million tokens (MTok). Fields: base input,
# output, cache writes, cache hits/reads. Sources: Anthropic and OpenAI
# pricing (verified 2026-09-17). Sonnet 5's launch rate of $2/$10 is now its
# standard price (the increase scheduled for 2026-09-01 was cancelled).
# OpenAI GPT-5.6 has no separate cache-write charge (cache writes are billed
# as normal input), so cache_write mirrors input for those entries.
_PRICING_PER_MTOK = {
    "claude-opus-5": {
        "input": 5.0, "output": 25.0, "cache_write": 6.25, "cache_read": 0.50,
    },
    "claude-sonnet-5": {
        "input": 2.0, "output": 10.0, "cache_write": 2.50, "cache_read": 0.20,
    },
    "gpt-5.6-sol": {
        "input": 5.0, "output": 30.0, "cache_write": 5.0, "cache_read": 0.50,
    },
    "gpt-5.6-terra": {
        "input": 2.50, "output": 15.0, "cache_write": 2.50, "cache_read": 0.25,
    },
    "gpt-5.6-luna": {
        "input": 1.0, "output": 6.0, "cache_write": 1.0, "cache_read": 0.10,
    },
}


def _model_pricing(model: str) -> dict | None:
    """Return the price table for a model, or None if it isn't on file."""
    return _PRICING_PER_MTOK.get(model)


# Prompts live in an external YAML file next to this script so they can be
# edited without touching the code.
PROMPTS_FILENAME = "term_extractor_prompts.yaml"

# Canonical column order for both the JSON records and the mirrored XLSX.
GLOSSARY_COLUMNS = [
    "source_term", "target_term", "term_type", "ngram", "seg_count",
    "raw_freq", "sources", "definition",
    "example_1", "target_example_1",
    "example_2", "target_example_2",
    "example_3", "target_example_3",
    "notes",
]
# For unspaced languages (Chinese, Japanese) the length column holds the
# term's CHARACTER count and is written out under this name instead of
# "ngram". Internally the field is always "ngram"; the rename happens only at
# output.
LENGTH_COLUMN_CHARS = "characters"


def _term_char_length(term: str) -> int:
    """Character length of a term, ignoring whitespace (so a mixed term like
    "USB ケーブル" counts its 7 visible characters)."""
    return sum(1 for ch in term if not ch.isspace())


def _output_columns(by_chars: bool) -> list[str]:
    """The public column list, with "ngram" renamed for character mode."""
    if not by_chars:
        return list(GLOSSARY_COLUMNS)
    return [LENGTH_COLUMN_CHARS if c == "ngram" else c
            for c in GLOSSARY_COLUMNS]

# N-gram lengths the GUI offers. Candidates longer than max(NGRAM_CHOICES)
# can never be extracted, which the run log distinguishes from lengths the
# user simply left unticked.
NGRAM_CHOICES = (1, 2, 3, 4, 5)

# Minimum segment length (in whitespace-split tokens) to process.
# Set from the GUI; this is the default. Low value keeps short but valid
# terms from lists/tables while the noise filter still drops bare single words
# and punctuation-only segments.
MIN_SEGMENT_TOKENS = 2

# --- spaCy language / model configuration ---
# Folder where extracted model directories live (e.g. it_core_news_lg-3.8.0).
# Override with the TERM_EXTRACTOR_MODELS_DIR environment variable; otherwise
# a per-user cache folder is used:
#   Windows : %LOCALAPPDATA%\term-extractor\spacy_models
#   macOS   : ~/Library/Caches/term-extractor/spacy_models
#   Linux   : $XDG_CACHE_HOME/term-extractor/spacy_models (~/.cache by default)
def _default_models_dir() -> Path:
    override = os.environ.get("TERM_EXTRACTOR_MODELS_DIR")
    if override:
        return Path(override).expanduser()
    if sys.platform == "win32":
        base = Path(os.environ.get("LOCALAPPDATA")
                    or Path.home() / "AppData" / "Local")
    elif sys.platform == "darwin":
        base = Path.home() / "Library" / "Caches"
    else:
        base = Path(os.environ.get("XDG_CACHE_HOME")
                    or Path.home() / ".cache")
    return base / "term-extractor" / "spacy_models"


SPACY_MODELS_DIR = _default_models_dir()

# Base download URL for spaCy model release wheels.
SPACY_MODEL_DOWNLOAD_BASE = "https://github.com/explosion/spacy-models/releases/download"

# Language name -> lg pipeline package name.
# English and Chinese use the "web" genre; all others use "news".
# (All 24 entries verified to have a large 'lg' model.)
LANGUAGE_MODELS = {
    "Catalan": "ca_core_news_lg",
    "Chinese": "zh_core_web_lg",
    "Croatian": "hr_core_news_lg",
    "Danish": "da_core_news_lg",
    "Dutch": "nl_core_news_lg",
    "English": "en_core_web_lg",
    "Finnish": "fi_core_news_lg",
    "French": "fr_core_news_lg",
    "German": "de_core_news_lg",
    "Greek": "el_core_news_lg",
    "Italian": "it_core_news_lg",
    "Japanese": "ja_core_news_lg",
    "Korean": "ko_core_news_lg",
    "Lithuanian": "lt_core_news_lg",
    "Macedonian": "mk_core_news_lg",
    "Norwegian Bokmål": "nb_core_news_lg",
    "Polish": "pl_core_news_lg",
    "Portuguese": "pt_core_news_lg",
    "Romanian": "ro_core_news_lg",
    "Russian": "ru_core_news_lg",
    "Slovenian": "sl_core_news_lg",
    "Spanish": "es_core_news_lg",
    "Swedish": "sv_core_news_lg",
    "Ukrainian": "uk_core_news_lg",
}

DEFAULT_LANGUAGE = "English"

# Languages written without spaces between words (GUI name -> spaCy code).
# Whitespace-based token counts are meaningless for them, so the
# Min. segment length filter is not applied: an unspaced sentence would count
# as one "token", and even a one-character segment can hold a real term.
UNSPACED_LANGUAGES = {"Chinese": "zh", "Japanese": "ja"}

# Accepted TMX srclang codes per language (ISO 639-1 base + common regional
# variants). Values are compared after stripping whitespace and lowercasing.
# The base code always matches any regional variant via prefix (e.g. "it-ch"),
# so this set mainly documents/anchors the expected forms.
LANGUAGE_SRC_CODES = {
    "Catalan": {"ca"},
    "Chinese": {"zh"},
    "Croatian": {"hr"},
    "Danish": {"da"},
    "Dutch": {"nl"},
    "English": {"en"},
    "Finnish": {"fi"},
    "French": {"fr"},
    "German": {"de"},
    "Greek": {"el"},
    "Italian": {"it"},
    "Japanese": {"ja"},
    "Korean": {"ko"},
    "Lithuanian": {"lt"},
    "Macedonian": {"mk"},
    "Norwegian Bokmål": {"nb", "no"},
    "Polish": {"pl"},
    "Portuguese": {"pt"},
    "Romanian": {"ro"},
    "Russian": {"ru"},
    "Slovenian": {"sl"},
    "Spanish": {"es"},
    "Swedish": {"sv"},
    "Ukrainian": {"uk"},
}

# Supported / recognised extensions
SUPPORTED_EXTS = {".tmx", ".pdf", ".docx"}
LEGACY_DOC_EXT = ".doc"


# ---------------------------------------------------------------------------
# SEGMENT MODEL
# ---------------------------------------------------------------------------

@dataclass
class Segment:
    """A unit of source text plus the file type it came from. For TMX segments
    the aligned target-language text is also kept; for monolingual sources
    (pdf, docx) 'target' stays empty."""
    text: str
    source: str   # "tmx", "tmx (source only[, AI-translated])", "pdf", "docx"
    target: str = ""   # aligned target text, TMX only


# ---------------------------------------------------------------------------
# LOGGING (routes print-style messages to GUI or console)
# ---------------------------------------------------------------------------

class Logger:
    """Routes log messages to a callback (GUI) or falls back to stdout,
    and keeps a buffer of every message so the run can be saved to a file."""
    def __init__(self, sink=None, update_sink=None, error_sink=None):
        self.sink = sink
        self.update_sink = update_sink
        self.error_sink = error_sink
        self.buffer: list[str] = []
        self._last_was_update = False

    @staticmethod
    def _stamp(msg: str) -> str:
        return f"[{datetime.now():%H:%M:%S}] {msg}"

    def log(self, msg: str):
        out = self._stamp(msg)
        self._last_was_update = False
        self.buffer.append(out)
        if self.sink:
            self.sink(out)
        else:
            print(out)

    def log_error(self, msg: str):
        """Log a prominent error/warning message (rendered in red in the GUI).
        Falls back to the normal sink/stdout when no error sink is set."""
        out = self._stamp(msg)
        self._last_was_update = False
        self.buffer.append(out)
        if self.error_sink:
            self.error_sink(out)
        elif self.sink:
            self.sink(out)
        else:
            print(out)

    def update(self, msg: str):
        """Update the last log line in place (e.g. a progress counter)."""
        out = self._stamp(msg)
        # Collapse consecutive update lines in the buffer so the saved log
        # keeps only the final counter value rather than every tick.
        if self.buffer and self._last_was_update:
            self.buffer[-1] = out
        else:
            self.buffer.append(out)
        self._last_was_update = True
        if self.update_sink:
            self.update_sink(out)
        else:
            print("\r" + out, end="", flush=True)

    def save_to_file(self, path: str):
        """Write all buffered log lines to a text file."""
        with open(path, "w", encoding="utf-8") as f:
            f.write("\n".join(self.buffer) + "\n")


# ---------------------------------------------------------------------------
# STEP 1: FILE DISCOVERY
# ---------------------------------------------------------------------------

def discover_files(folder: str, logger: Logger,
                   recursive: bool = True) -> dict[str, list[Path]]:
    """
    Find supported files, grouped by extension. Scans subfolders when
    `recursive` is True, otherwise only the top-level folder.
    Logs how many of each type were found, and warns about skipped .doc files.
    """
    by_ext: dict[str, list[Path]] = defaultdict(list)
    legacy_doc = []

    entries = Path(folder).rglob("*") if recursive else Path(folder).glob("*")
    for p in sorted(entries):
        if not p.is_file():
            continue
        ext = p.suffix.lower()
        if ext in SUPPORTED_EXTS:
            by_ext[ext].append(p)
        elif ext == LEGACY_DOC_EXT:
            legacy_doc.append(p)

    scope = "recursively" if recursive else "top-level only"
    logger.log(f"[discover] Scanned {scope}. TMX: {len(by_ext['.tmx'])}, "
               f"PDF: {len(by_ext['.pdf'])}, DOCX: {len(by_ext['.docx'])}")
    if legacy_doc:
        logger.log(f"[discover] Skipping {len(legacy_doc)} legacy .doc file(s) "
                   f"(unsupported format).")
    return by_ext


# ---------------------------------------------------------------------------
# STEP 1a: TMX PARSING
# ---------------------------------------------------------------------------

def _norm_lang(code: str) -> str:
    """Normalise a language code: strip surrounding whitespace, lowercase."""
    return (code or "").strip().lower()


def read_tmx_srclang(path: Path) -> str | None:
    """
    Return the normalised srclang declared in a TMX <header>, or None if the
    file can't be read or declares no srclang.
    """
    try:
        tree = ET.parse(path)
        header = tree.getroot().find("header")
        if header is None:
            return None
        return _norm_lang(header.get("srclang", "")) or None
    except Exception:
        return None


def srclang_matches_language(srclang: str | None, language: str) -> bool:
    """
    True if a normalised srclang belongs to the selected UI language.
    Matches the base ISO code by prefix, so 'it', 'it-it', 'it-ch' all match
    Italian. '*all*', missing, or a different language returns False.
    """
    if not srclang or srclang == "*all*":
        return False
    accepted = LANGUAGE_SRC_CODES.get(language, set())
    for code in accepted:
        if srclang == code or srclang.startswith(code + "-"):
            return True
    return False


def check_tmx_source_languages(paths: list[Path], language: str):
    """
    Split TMX files into (compatible, problematic) based on their srclang vs
    the selected language. `problematic` is a list of (path, reason) tuples.
    """
    compatible: list[Path] = []
    problematic: list[tuple[Path, str]] = []
    for path in paths:
        srclang = read_tmx_srclang(path)
        if srclang is None:
            problematic.append((path, "no srclang in header"))
        elif srclang == "*all*":
            problematic.append((path, "srclang='*all*'"))
        elif srclang_matches_language(srclang, language):
            compatible.append(path)
        else:
            problematic.append((path, f"srclang='{srclang}'"))
    return compatible, problematic


def collect_tmx_target_codes(path: Path) -> list[str]:
    """
    Return the sorted list of distinct target-side language codes present in a
    TMX file: every TUV xml:lang that does NOT belong to the file's own source
    language. Codes are returned in their original (normalised) form, e.g.
    ['de-de', 'it-ch', 'it-it']. Empty list if the file can't be read or has no
    distinct target languages.
    """
    srclang = read_tmx_srclang(path)
    src_root = srclang.split("-")[0] if srclang else ""
    codes: set[str] = set()
    try:
        tree = ET.parse(path)
        root = tree.getroot()
        for tu in root.iter("tu"):
            for tuv in tu.findall("tuv"):
                lang = _norm_lang(
                    tuv.get("{http://www.w3.org/XML/1998/namespace}lang")
                    or tuv.get("lang", ""))
                if not lang:
                    continue
                is_source = bool(srclang) and (
                    lang == srclang
                    or lang.startswith(src_root + "-")
                    or lang.split("-")[0] == src_root)
                if not is_source:
                    codes.add(lang)
    except Exception:
        return []
    return sorted(codes)


def resolve_tmx_target_codes(codes_by_file: dict[str, list[str]],
                             target_language: str, prompt_template: str,
                             client, model: str, logger: Logger,
                             usage_totals: dict | None = None) -> dict:
    """
    Ask the model which of each TMX file's target-side language codes match the
    user's requested target language. Sends one batched request: the target
    string plus a per-file list of codes. Returns a dict:

        { filename: {"matches": [codes...]} }

    where 'matches' is the subset of that file's codes the model judges
    compatible with the target language (possibly empty, possibly several).
    The model is limited to choosing from each file's actual codes; any code it
    returns that wasn't offered is discarded (defensive). Token usage is
    accumulated into usage_totals like every other AI call.

    Raises on a failed/unparseable call so the caller can decide how to handle
    it (this runs at pre-flight, before heavy work).
    """
    # Build a compact payload keyed by filename (basename, for a short prompt).
    payload = {Path(f).name: c for f, c in codes_by_file.items()}
    prompt = _fill_placeholders(prompt_template, {
        "TARGET_LANGUAGE": target_language,
        "FILES_JSON": json.dumps(payload, ensure_ascii=False, indent=2),
    })

    result, usage, truncated, _raw = _call_ai_json(client, prompt, model, logger)
    if usage_totals is not None:
        _accumulate_usage(usage_totals, usage)
    if truncated or not isinstance(result, dict):
        raise ValueError("Language-resolution response was unusable "
                         "(truncated or not an object).")

    files = result.get("files")
    if not isinstance(files, dict):
        raise ValueError("Language-resolution response had no 'files' object.")

    # Map the model's basename-keyed answer back to the original file keys,
    # keeping only codes the file actually offered.
    out: dict = {}
    name_to_key = {Path(f).name: f for f in codes_by_file}
    for name, key in name_to_key.items():
        offered = set(codes_by_file[key])
        entry = files.get(name) or {}
        matches = entry.get("matches") if isinstance(entry, dict) else None
        if not isinstance(matches, list):
            matches = []
        kept = [c for c in matches if c in offered]
        out[key] = {"matches": kept}
    return out


def parse_tmx_files(paths: list[Path], logger: Logger,
                    target_code_by_file: dict[str, str | None] | None = None,
                    will_translate: bool = False) -> list[Segment]:
    """
    Parse TMX files -> one Segment per source-language TU, carrying the aligned
    target text.

    The source side is chosen per file from its own header srclang.

    Target selection depends on target_code_by_file (keyed by str(path)):
      - If a file maps to a specific code, the target is the TUV whose language
        matches THAT EXACT code (region-specific), so multilingual TMs don't mix
        locales. TUs lacking that code get an empty target. Such segments are
        tagged source "tmx".
      - If a file maps to None (no compatible target, or the user chose not to
        use its targets), every segment gets an empty target and is tagged
        source-only. The exact tag depends on will_translate: when the AI
        translation step will run, those terms WILL be AI-translated, so the
        tag is TAG_TMX_SOURCE_ONLY_AI; otherwise nothing is translated and the
        plain TAG_TMX_SOURCE_ONLY is used.
      - If target_code_by_file is None entirely (no resolution performed at
        all), the first non-source TUV is used as the target and segments are
        tagged "tmx".
    """
    segments: list[Segment] = []
    total = len(paths)
    source_only_tag = (TAG_TMX_SOURCE_ONLY_AI if will_translate
                       else TAG_TMX_SOURCE_ONLY)
    for i, path in enumerate(paths, start=1):
        if total:
            logger.update(f"[tmx] Reading {i}/{total} TMX file(s)...")
        srclang = read_tmx_srclang(path)   # already validated upstream
        src_root = srclang.split("-")[0] if srclang else ""

        # Decide this file's target handling.
        if target_code_by_file is None:
            chosen_code = "*first*"       # default: first non-source TUV
            source_tag = "tmx"
        else:
            chosen_code = target_code_by_file.get(str(path))
            source_tag = "tmx" if chosen_code else source_only_tag

        try:
            tree = ET.parse(path)
            root = tree.getroot()
            for tu in root.iter("tu"):
                source_text = None
                target_text = ""
                for tuv in tu.findall("tuv"):
                    lang = _norm_lang(
                        tuv.get("{http://www.w3.org/XML/1998/namespace}lang")
                        or tuv.get("lang", ""))
                    seg = tuv.find("seg")
                    seg_text = seg.text.strip() if (seg is not None and seg.text) else ""
                    is_source = bool(srclang) and (
                        lang == srclang
                        or lang.startswith(src_root + "-")
                        or lang.split("-")[0] == src_root)
                    if is_source:
                        if source_text is None and seg_text:
                            source_text = seg_text
                    elif chosen_code and seg_text:
                        # Target selection. Normalise tags/whitespace but never
                        # noise-filter — we don't extract from the target.
                        if chosen_code == "*first*":
                            if not target_text:
                                target_text = clean_text(seg_text)
                        elif lang == chosen_code and not target_text:
                            target_text = clean_text(seg_text)
                if source_text:
                    segments.append(Segment(source_text, source_tag, target_text))
        except Exception as e:
            logger.log(f"\n[tmx]   ! Failed to read '{path.name}': {e}")
    if total:
        with_target = sum(1 for s in segments if s.target)
        source_only = sum(1 for s in segments
                          if s.source in (TAG_TMX_SOURCE_ONLY,
                                          TAG_TMX_SOURCE_ONLY_AI))
        logger.log(f"[tmx] Extracted {len(segments)} TU segment(s) "
                   f"({with_target} with aligned target"
                   + (f", {source_only} source-only" if source_only else "")
                   + ").")
    return segments


# ---------------------------------------------------------------------------
# STEP 1b: PDF PARSING
# ---------------------------------------------------------------------------

def parse_pdf_files(paths: list[Path], nlp, logger: Logger) -> list[Segment]:
    """
    Extract selectable text from PDFs, page by page (image-only pages skipped),
    then split into sentence Segments.
    """
    if not paths:
        return []
    import pdfplumber

    text_blocks: list[str] = []
    skipped_pages = 0
    processed_pages = 0
    total = len(paths)

    for i, path in enumerate(paths, start=1):
        logger.update(f"[pdf] Reading {i}/{total} PDF file(s)...")
        try:
            with pdfplumber.open(path) as pdf:
                for page in pdf.pages:
                    text = page.extract_text() or ""
                    if not text.strip():
                        skipped_pages += 1
                        continue
                    text_blocks.append(text)
                    processed_pages += 1
        except Exception as e:
            logger.log(f"\n[pdf]   ! Failed to read '{path.name}': {e}")

    logger.log(f"[pdf] Processed {processed_pages} text page(s), "
               f"skipped {skipped_pages} image-only page(s).")

    segments = _sentences_from_blocks(text_blocks, "pdf", nlp)
    logger.log(f"[pdf] Produced {len(segments)} sentence segment(s).")
    return segments


# ---------------------------------------------------------------------------
# STEP 1c: DOCX PARSING
# ---------------------------------------------------------------------------

def parse_docx_files(paths: list[Path], nlp, logger: Logger) -> list[Segment]:
    """
    Extract text from DOCX paragraphs and table cells,
    then split into sentence Segments.
    """
    if not paths:
        return []
    import docx  # python-docx

    text_blocks: list[str] = []
    total = len(paths)

    for i, path in enumerate(paths, start=1):
        logger.update(f"[docx] Reading {i}/{total} DOCX file(s)...")
        try:
            document = docx.Document(str(path))
            # Paragraphs
            for para in document.paragraphs:
                if para.text and para.text.strip():
                    text_blocks.append(para.text)
            # Table cells
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text and cell.text.strip():
                            text_blocks.append(cell.text)
        except Exception as e:
            logger.log(f"\n[docx]   ! Failed to read '{path.name}': {e}")

    logger.log(f"[docx] Extracted {len(text_blocks)} text block(s).")

    segments = _sentences_from_blocks(text_blocks, "docx", nlp)
    logger.log(f"[docx] Produced {len(segments)} sentence segment(s).")
    return segments


def _sentences_from_blocks(blocks: list[str], source: str, nlp) -> list[Segment]:
    """Split a list of text blocks into sentence Segments tagged with `source`."""
    segments: list[Segment] = []
    for doc in nlp.pipe(blocks, batch_size=32):
        for sent in doc.sents:
            s = sent.text.strip()
            if s:
                segments.append(Segment(s, source))
    return segments


# ---------------------------------------------------------------------------
# STEP 2: CLEANING
# ---------------------------------------------------------------------------

# (reason, pattern) pairs — the reason is reported in the run log so it is
# clear why a segment was discarded.
_NOISE_PATTERNS = [
    ("digits/punctuation only", re.compile(r"^[\d\s\W]+$")),
    ("code-like",               re.compile(r"^[A-Z0-9\-_]+$")),
    ("contains a URL",          re.compile(r"https?://\S+")),
    ("contains an email",       re.compile(r"\S+@\S+\.\S+")),
    ("UI string",               re.compile(r"^(?:yes|no|ok|cancel)$", re.I)),
]

# Matches genuine inline XML/HTML tags only: '<', optional '/', a letter-led
# tag name (with optional namespace prefix like 'w:tab'), optional attributes,
# optional self-closing '/', '>'. This leaves prose-in-brackets and math like
# 'a < b' intact, while removing real tags such as <ph>, </seg>, <bpt i="1">,
# and <x/>. Malformed pseudo-tags are rare in real corpora and fine to lose.
_TAG_RE = re.compile(
    r"</?[A-Za-z][A-Za-z0-9]*(?::[A-Za-z][A-Za-z0-9]*)?(?:\s+[^<>]*?)?/?>")

def clean_text(text: str) -> str:
    """Remove inline XML/HTML tags, resolve HTML entities, normalize whitespace.
    Only genuine tag-shaped substrings are stripped (see _TAG_RE), so content
    like 'a < b' survives. Tags are removed first, then entities are unescaped
    to their real characters (so 'AT&amp;T' becomes 'AT&T', not 'AT T'), then
    whitespace is collapsed."""
    text = _TAG_RE.sub(" ", text)
    text = html.unescape(text)
    text = re.sub(r"\s+", " ", text).strip()
    return text

def is_noisy(text: str) -> str | None:
    for reason, pattern in _NOISE_PATTERNS:
        if pattern.search(text):
            return reason
    return None

def clean_segments(segments: list[Segment], logger: Logger) -> list[Segment]:
    """Clean text and drop noisy segments, preserving source tags. Reports a
    per-reason breakdown so it is clear what was discarded and why."""
    cleaned: list[Segment] = []
    dropped: Counter = Counter()
    for seg in segments:
        text = clean_text(seg.text)
        if not text:
            dropped["empty after tag/entity cleanup"] += 1
            continue
        reason = is_noisy(text)
        if reason:
            dropped[reason] += 1
            continue
        cleaned.append(Segment(text, seg.source, seg.target))
    total_dropped = sum(dropped.values())
    if total_dropped:
        detail = ", ".join(f"{r}: {n}" for r, n in dropped.most_common())
        logger.log(f"[clean] Dropped {total_dropped} of {len(segments)} "
                   f"segment(s) — {detail}.")
    logger.log(f"[clean] {len(cleaned)} segment(s) after cleaning.")
    return cleaned


# ---------------------------------------------------------------------------
# STEP 3: CANDIDATE EXTRACTION (spaCy)
# ---------------------------------------------------------------------------

def _find_model_dir(package: str) -> Path | None:
    """
    Look in SPACY_MODELS_DIR for a subfolder whose name starts with the package
    name (folders carry a version suffix, e.g. 'it_core_news_lg-3.8.0').
    Returns the path to the loadable model dir, or None if not present.
    """
    base = Path(SPACY_MODELS_DIR)
    if not base.is_dir():
        return None
    # Prefer an exact-prefix match with a version suffix; fall back to exact name.
    candidates = sorted(
        p for p in base.iterdir()
        if p.is_dir() and (p.name == package or p.name.startswith(package + "-"))
    )
    return candidates[-1] if candidates else None


def _download_model(package: str, logger: Logger) -> Path:
    """
    Download the spaCy model wheel for `package`, extract its inner versioned
    model directory into SPACY_MODELS_DIR, and return the extracted path.
    """
    import io
    import zipfile
    import urllib.request
    from spacy.cli.download import get_compatibility

    base = Path(SPACY_MODELS_DIR)
    base.mkdir(parents=True, exist_ok=True)

    # Resolve the exact model version compatible with the installed spaCy.
    compat = get_compatibility()
    if package not in compat:
        raise RuntimeError(f"No compatible version found for '{package}'.")
    version = compat[package][0]

    wheel_name = f"{package}-{version}-py3-none-any.whl"
    url = f"{SPACY_MODEL_DOWNLOAD_BASE}/{package}-{version}/{wheel_name}"

    logger.log(f"[model] '{package}' not found — downloading {package}-{version}...")
    logger.log(f"[model] This is a large file and may take a while.")

    with urllib.request.urlopen(url) as resp:
        data = resp.read()

    # The wheel contains: <package>/<package>-<version>/  (the loadable model dir)
    inner_prefix = f"{package}/{package}-{version}/"
    dst = base / f"{package}-{version}"
    if dst.exists():
        import shutil
        shutil.rmtree(dst)

    with zipfile.ZipFile(io.BytesIO(data)) as z:
        members = [n for n in z.namelist() if n.startswith(inner_prefix)]
        for n in members:
            rel = n[len(inner_prefix):]
            if not rel:
                continue
            target = dst / rel
            if n.endswith("/"):
                target.mkdir(parents=True, exist_ok=True)
            else:
                target.parent.mkdir(parents=True, exist_ok=True)
                with z.open(n) as src, open(target, "wb") as out:
                    out.write(src.read())

    logger.log(f"[model] Installed to '{dst}'.")
    return dst


def load_spacy_model(language: str, logger: Logger):
    """
    Resolve the lg model for the selected language, downloading it into
    SPACY_MODELS_DIR if not already present, then load it by path.
    """
    package = LANGUAGE_MODELS[language]
    model_dir = _find_model_dir(package)
    if model_dir is None:
        model_dir = _download_model(package, logger)
    logger.log(f"[model] Loading {package} ({language}) from '{model_dir}'...")
    return spacy.load(str(model_dir))

def _get_chunks(doc):
    """
    Return the document's noun chunks. Not every spaCy pipeline implements
    the noun_chunks iterator for its language (e.g. Russian raises
    NotImplementedError); in that case fall back to grouping consecutive
    ADJ/NOUN/PROPN runs, which works for any language.
    """
    try:
        return list(doc.noun_chunks)
    except (NotImplementedError, ValueError):
        return _fallback_chunks(doc)


def _fallback_chunks(doc):
    """Group consecutive ADJ/NOUN/PROPN token runs into spans."""
    chunks = []
    start = None
    for i, tok in enumerate(doc):
        if tok.pos_ in ("NOUN", "PROPN", "ADJ"):
            if start is None:
                start = i
        else:
            if start is not None:
                chunks.append(doc[start:i])
                start = None
    if start is not None:
        chunks.append(doc[start:len(doc)])
    return chunks


def is_valid_noun_chunk(chunk, custom_stopwords: set) -> bool:
    if chunk.root.pos_ not in ("NOUN", "PROPN"):
        return False
    tokens_lower = [t.text.lower() for t in chunk]
    if any(t in custom_stopwords for t in tokens_lower):
        return False
    if all(t.pos_ in ("DET", "PRON", "ADP") for t in chunk):
        return False
    return True


def is_junk_term(text: str, chunk) -> str | None:
    """
    Return a short reason string if `text` is extraction junk rather than a
    candidate term, or None if it looks like a real term. `chunk` is the
    parsed spaCy span the text came from; its POS tags are used for the
    measurement test.

    The rules are deliberately high-precision: a wrongly dropped term is
    invisible and unrecoverable, whereas junk that slips through is visible in
    the output and is caught by the AI cleaning step. Everything here is
    language-independent — the patterns are structural, content words are
    identified by POS tag rather than by word list, and only symbol units are
    listed (see TERM_JUNK_UNITS).

    Letters are matched in any script (Latin, Cyrillic, Greek, CJK, ...).
    Two rules are script-aware: "too short" is skipped for terms containing
    CJK characters (one or two characters is a normal term length there), and
    "code pattern" only counts alphabetic-script letters, since CJK characters
    next to digits are ordinary prose ("3号機"), not product codes.

    Rules, in order:
      too short        - 1-2 chars with no letter pair          "[A"
                         (not applied to CJK terms)
      no letters       - no letter in any script                "+33"
      long number      - 5+ digits ignoring . - / space         "18882686222"
      punctuation-led  - opens with a bracket/quote/pipe        "(inches"
      unbalanced ...   - brackets or quotes don't pair up       "code “US"
      code pattern     - 2+ letters next to 2+ digits, either   "GU14"
                         order; skipped when the letters are a
                         unit, so "10ml syringe" survives;
                         alphabetic-script letters only
      bare unit        - the whole term is a unit               "cm"
      measurement only - no content word that isn't a unit      "less than 8 cm"
    """
    t = (text or "").strip()
    if not t:
        return "empty"
    if (len(t) <= 2 and not _CJK_RE.search(t)
            and not _ALPHA_PAIR_RE.search(t)):
        return "too short"
    if not _ANY_LETTER_RE.search(t):
        return "no letters"
    # Long digit runs (phone numbers, document codes), ignoring separators.
    if re.search(r"\d{5,}", re.sub(r"[.\-/\s]", "", t)):
        return "long number"
    if t[0] in TERM_JUNK_OPENERS:
        return "punctuation-led"
    for opener, closer in (("(", ")"), ("[", "]"), ("{", "}")):
        if t.count(opener) != t.count(closer):
            return "unbalanced brackets"
    if (t.count('"') + t.count("“") + t.count("”")) % 2:
        return "unbalanced quotes"
    # Alphanumeric code, either order. A digits-then-letters match whose
    # letters are a unit is a glued measurement ("10ml"), not a code. Only
    # alphabetic-script letters count (see _CODE_PATTERN_RE).
    for m in _CODE_PATTERN_RE.finditer(t):
        letters = m.group(1) or m.group(4)
        if m.group(3) and letters.lower() in TERM_JUNK_UNITS:
            continue
        return "code pattern"
    if t.lower().strip(".,") in TERM_JUNK_UNITS:
        return "bare unit"

    # Measurement-only: does the chunk contain any content NOUN that isn't a
    # unit? Only NOUN/PROPN count — an adjective like "less" in "less than
    # 8 cm" is a measurement qualifier, not content. is_valid_noun_chunk has
    # already guaranteed the chunk is noun-headed, so if every noun in it is a
    # unit, the whole chunk is a measurement. POS tags keep this
    # language-independent.
    has_content = any(
        tok.pos_ in ("NOUN", "PROPN")
        and tok.text.lower().strip(".,") not in TERM_JUNK_UNITS
        for tok in chunk)
    if not has_content and re.search(r"\d", t):
        return "measurement only"
    return None


def _clean_chunk_span(chunk):
    """
    Strip leading determiners/pronouns from a chunk (e.g. "the safety valve"
    -> "safety valve") and return (clean_text, start_char, end_char), where
    the offsets index into the original segment text (chunk.doc.text) so the
    term can be bolded inside its example sentence. Returns None if nothing
    valid remains after cleaning.
    """
    tokens = list(chunk)
    while tokens and tokens[0].pos_ in ("DET", "PRON"):
        tokens = tokens[1:]
    if not tokens:
        return None

    start = tokens[0].idx
    last = tokens[-1]
    end = last.idx + len(last.text)
    text = chunk.doc.text[start:end]

    # Trim residual whitespace at the edges, adjusting the offsets.
    lstripped = text.lstrip()
    start += len(text) - len(lstripped)
    stripped = lstripped.rstrip()
    end -= len(lstripped) - len(stripped)

    # Single characters are dropped, except in CJK scripts where one
    # character is a normal term length ("弁", "軸").
    if len(stripped) < 2 and not _CJK_RE.search(stripped):
        return None
    return stripped, start, end


def _chunk_root_pos(chunk):
    """POS of the chunk's head token, falling back to the most common
    content POS among its tokens if the head isn't reliable."""
    try:
        if chunk.root.pos_:
            return chunk.root.pos_
    except Exception:
        pass
    counts = Counter(t.pos_ for t in chunk if t.pos_ not in ("DET", "PRON", "ADP", ""))
    return counts.most_common(1)[0][0] if counts else None


def _term_type_label(pos):
    if pos in ("NOUN", "PROPN"):
        return "Noun"
    if pos == "ADJ":
        return "Adjective"
    if pos == "VERB":
        return "Verb"
    return "Other"


def _make_bold_example(text, start, end):
    """Return the segment text with the term (by offsets) wrapped in **bold**."""
    return f"{text[:start]}**{text[start:end]}**{text[end:]}"


def _bold_forms_in_text(text: str, forms: list[str]) -> str:
    """
    Wrap every verbatim occurrence of any form in `forms` with **bold** in
    `text`. Matching is case-sensitive and literal (the model returns forms
    copied verbatim from the segment). Longer forms are applied first so a
    shorter form doesn't match inside a longer one, and already-bolded spans
    are never re-wrapped. Forms not found are simply skipped (safe fallback:
    the text is returned unchanged for those).
    """
    if not text or not forms:
        return text
    # Distinct, non-empty forms, longest first.
    ordered = sorted({f for f in forms if f and f.strip()},
                     key=len, reverse=True)
    for form in ordered:
        # Split on any existing bold spans so we never bold inside one, and
        # never double-wrap.
        parts = re.split(r"(\*\*.*?\*\*)", text)
        for idx, part in enumerate(parts):
            if part.startswith("**") and part.endswith("**"):
                continue   # leave existing bold spans untouched
            if form in part:
                # Literal replace of every occurrence in this non-bold part.
                parts[idx] = part.replace(form, f"**{form}**")
        text = "".join(parts)
    return text

def extract_candidates_spacy(segments: list[Segment], nlp,
                             selected_ngrams: set[int], custom_stopwords: set,
                             min_segment_tokens: int,
                             logger: Logger,
                             junk_filter: bool = True) -> dict[str, dict]:
    """
    Extract noun chunk candidates per segment using spaCy.

    For each candidate (keyed by its cleaned, lowercased form) tracks:
      - forms      : Counter of the actual surface spellings seen
      - pos_counts : Counter of head POS, to classify the term type
      - examples   : a pool of example dicts (see below), from which the
                     final examples are chosen later in build_records
    Leading determiners/pronouns are stripped ("the safety valve" ->
    "safety valve"); languages without noun_chunks use a fallback.

    Each example dict is:
        {"source": "<segment text, term in **bold**>",
         "target": "<aligned target segment, TMX only, else ''>",
         "origin": "tmx" | "pdf" | "docx"}

    We keep up to EXAMPLE_POOL_MAX examples per term (more than MAX_EXAMPLES)
    so the later selection can prefer TMX-aligned, shortest-first examples.

    Returns: { key: {"forms": Counter, "pos_counts": Counter,
                      "examples": [ {source,target,origin}, ... ]} }
    """
    candidates: dict[str, dict] = defaultdict(lambda: {
        "forms": Counter(),
        "pos_counts": Counter(),
        "examples": [],
    })
    junk_counts: Counter = Counter()

    # Keep the Segment alongside each doc so we can capture the aligned target
    # and origin per example (as_tuples carries the context object through).
    unspaced = nlp.lang in UNSPACED_LANGUAGES.values()
    if unspaced:
        work = [(s.text, s) for s in segments]
        logger.log("[spacy] Min. segment length not applied (language is "
                   "written without spaces between words).")
    else:
        work = [(s.text, s) for s in segments
                if len(s.text.split()) >= min_segment_tokens]
    too_short = len(segments) - len(work)
    if too_short:
        logger.log(f"[spacy] Skipped {too_short} segment(s) shorter than "
                   f"{min_segment_tokens} token(s); {len(work)} segment(s) "
                   f"parsed.")

    for doc, seg in nlp.pipe(work, batch_size=64, as_tuples=True):
        seg_text = doc.text
        seen_in_seg = set()
        for chunk in _get_chunks(doc):
            if not is_valid_noun_chunk(chunk, custom_stopwords):
                continue
            cleaned = _clean_chunk_span(chunk)
            if not cleaned:
                continue
            text, start, end = cleaned

            # Drop extraction junk (codes, phone numbers, bare measurements,
            # truncated fragments) before it enters the pipeline at all.
            if junk_filter:
                junk_reason = is_junk_term(text, chunk)
                if junk_reason:
                    junk_counts[junk_reason] += 1
                    continue

            # n-gram length is measured on the cleaned form (after determiner
            # stripping), so "the safety valve" counts as a 2-gram.
            # Unspaced languages (Chinese, Japanese) have no word count to
            # filter on, so every length is kept.
            if not unspaced and len(text.split()) not in selected_ngrams:
                continue

            key = text.lower()
            entry = candidates[key]
            entry["forms"][text] += 1
            entry["pos_counts"][_chunk_root_pos(chunk) or "X"] += 1
            if key not in seen_in_seg:
                seen_in_seg.add(key)
                if len(entry["examples"]) < EXAMPLE_POOL_MAX:
                    entry["examples"].append({
                        "source": _make_bold_example(seg_text, start, end),
                        "target": seg.target,
                        "origin": seg.source,
                    })

    if not junk_filter:
        logger.log("[spacy] Junk filter is OFF — raw candidates kept.")
    if junk_counts:
        detail = ", ".join(f"{reason}: {n}"
                           for reason, n in junk_counts.most_common())
        logger.log(f"[spacy] Junk filter removed {sum(junk_counts.values())} "
                   f"candidate occurrence(s) — {detail}.")
    logger.log(f"[spacy] {len(candidates)} unique candidate term(s) extracted.")
    return dict(candidates)

# ---------------------------------------------------------------------------
# AI STEP 0: PROMPT LOADING + API CLIENT
# ---------------------------------------------------------------------------

def load_prompts(logger: Logger) -> dict:
    """
    Load the cleaning/translation prompts from the YAML file that sits next to
    this script. Raise a clear error if the file is missing, unreadable, or
    doesn't contain both required keys.
    """
    import yaml  # local import so pure-NLP runs don't need pyyaml installed

    prompts_path = Path(__file__).resolve().parent / PROMPTS_FILENAME
    if not prompts_path.is_file():
        raise FileNotFoundError(
            f"Prompt file not found: {prompts_path}. It must sit next to the "
            f"script. AI steps cannot run without it.")

    try:
        with open(prompts_path, "r", encoding="utf-8") as fh:
            data = yaml.safe_load(fh)
    except Exception as e:
        raise ValueError(f"Could not parse '{PROMPTS_FILENAME}': {e}") from e

    if not isinstance(data, dict):
        raise ValueError(f"'{PROMPTS_FILENAME}' must contain a top-level mapping.")
    missing = [k for k in ("cleaning", "translation", "tmx_translation",
                           "tmx_language_resolution")
               if not data.get(k)]
    if missing:
        raise ValueError(
            f"'{PROMPTS_FILENAME}' is missing required prompt key(s): "
            f"{', '.join(missing)}.")

    logger.log(f"[ai] Prompts loaded from '{prompts_path.name}'.")
    return data


def make_ai_client(api_key: str, provider: str):
    """Construct an API client for the given provider ('anthropic' or
    'openai'). SDKs are imported lazily so pure-NLP runs need neither package,
    and an Anthropic-only run doesn't require the OpenAI package (or vice
    versa).

    Both SDKs retry transient failures (connection errors, 429 rate limits,
    5xx) internally with backoff; AI_SDK_MAX_RETRIES raises that count above
    the library default. Note this covers transport-level failures only — a
    successful response containing malformed JSON is not retried (it is
    reported in the log instead).

    Returns a (provider, client) tuple.
    """
    if provider == "openai":
        from openai import OpenAI
        return ("openai", OpenAI(api_key=api_key,
                                 max_retries=AI_SDK_MAX_RETRIES))
    from anthropic import Anthropic
    return ("anthropic", Anthropic(api_key=api_key,
                                   max_retries=AI_SDK_MAX_RETRIES))


def _fill_placeholders(template: str, mapping: dict[str, str]) -> str:
    """Replace {{KEY}} placeholders with values. Simple, dependency-free."""
    out = template
    for key, val in mapping.items():
        out = out.replace("{{" + key + "}}", val)
    return out


def _new_usage_totals() -> dict:
    """A per-run token accumulator."""
    return {"input": 0, "output": 0,
            "cache_creation": 0, "cache_read": 0, "calls": 0}


def _accumulate_usage(totals: dict, usage: dict | None) -> None:
    """Add one response's normalized usage into the running per-run totals.
    'usage' is a provider-neutral dict from _normalize_usage (or None). The
    'input' bucket is the uncached input; cache reads/writes are separate."""
    if usage is None:
        return
    totals["input"] += usage.get("input", 0)
    totals["output"] += usage.get("output", 0)
    totals["cache_creation"] += usage.get("cache_creation", 0)
    totals["cache_read"] += usage.get("cache_read", 0)
    totals["calls"] += 1


def _normalize_usage(provider: str, usage) -> dict | None:
    """Map a provider's usage object onto a common shape:
        {input (uncached), output, cache_creation, cache_read}

    Anthropic: input_tokens already excludes cached tokens; cache creation and
    read are separate fields.
    OpenAI: prompt_tokens INCLUDES cached tokens (cached_tokens is nested under
    prompt_tokens_details), so uncached input = prompt_tokens - cached; OpenAI
    has no separate cache-write charge, so cache_creation is 0.
    """
    if usage is None:
        return None
    if provider == "openai":
        prompt = getattr(usage, "prompt_tokens", 0) or 0
        completion = getattr(usage, "completion_tokens", 0) or 0
        details = getattr(usage, "prompt_tokens_details", None)
        cached = 0
        if details is not None:
            cached = (getattr(details, "cached_tokens", 0)
                      or (details.get("cached_tokens", 0)
                          if isinstance(details, dict) else 0) or 0)
        return {
            "input": max(0, prompt - cached),
            "output": completion,
            "cache_creation": 0,
            "cache_read": cached,
        }
    # anthropic
    return {
        "input": getattr(usage, "input_tokens", 0) or 0,
        "output": getattr(usage, "output_tokens", 0) or 0,
        "cache_creation": getattr(usage, "cache_creation_input_tokens", 0) or 0,
        "cache_read": getattr(usage, "cache_read_input_tokens", 0) or 0,
    }


def _salvage_json_array(text: str, array_key: str) -> list | None:
    """
    Recover as many COMPLETE objects as possible from a truncated JSON response
    of the form {"<array_key>": [ {...}, {...}, <cut off> .

    Strategy: find the array, then walk it accumulating complete top-level
    objects (balanced braces, respecting strings/escapes) until the text runs
    out or stops being parseable. Returns the list of parsed objects, or None
    if nothing usable could be recovered.
    """
    # Locate the start of the array belonging to array_key.
    m = re.search(r'"' + re.escape(array_key) + r'"\s*:\s*\[', text)
    if not m:
        return None
    i = m.end()               # position just after the '['
    objects = []
    depth = 0
    start = None
    in_str = False
    escape = False
    for pos in range(i, len(text)):
        ch = text[pos]
        if in_str:
            if escape:
                escape = False
            elif ch == "\\":
                escape = True
            elif ch == '"':
                in_str = False
            continue
        if ch == '"':
            in_str = True
        elif ch == "{":
            if depth == 0:
                start = pos
            depth += 1
        elif ch == "}":
            depth -= 1
            if depth == 0 and start is not None:
                chunk = text[start:pos + 1]
                try:
                    objects.append(json.loads(chunk))
                except json.JSONDecodeError:
                    break   # give up at the first unparseable complete-looking obj
                start = None
        elif ch == "]" and depth == 0:
            break   # clean end of array
    return objects or None


def _call_ai_json(client_tuple, prompt: str, model: str, logger: Logger):
    """
    Send a single-user-message request and parse the response as JSON.
    Works with either provider (client_tuple is (provider, client)). Returns
    (parsed_dict_or_None, normalized_usage_dict, truncated_bool, raw_text).

    On truncation the response is NOT raised on here; the flag and raw text are
    returned so the caller can decide whether to salvage partial results
    (cleaning) or skip the batch (translation). A hard ValueError is raised
    only for a NON-truncated response that fails to parse.
    """
    provider, client = client_tuple

    if provider == "openai":
        base_kwargs = {
            "model": model,
            "messages": [{"role": "user", "content": prompt}],
        }
        # Newer OpenAI SDKs/models use 'max_completion_tokens'; older ones only
        # accept 'max_tokens'. Try the modern name, fall back on a TypeError
        # about the unexpected keyword. Likewise, request JSON output but drop
        # response_format if the installed SDK/model rejects it.
        def _openai_create(token_kw, with_json):
            kw = dict(base_kwargs)
            kw[token_kw] = AI_MAX_TOKENS
            if with_json:
                kw["response_format"] = {"type": "json_object"}
            return client.chat.completions.create(**kw)

        attempts = [
            ("max_completion_tokens", True),
            ("max_tokens", True),
            ("max_completion_tokens", False),
            ("max_tokens", False),
        ]
        resp = None
        last_err = None
        for token_kw, with_json in attempts:
            try:
                resp = _openai_create(token_kw, with_json)
                break
            except TypeError as e:
                # Unexpected/!unsupported keyword for this SDK version; try next.
                last_err = e
                continue
        if resp is None:
            raise last_err
        choice = resp.choices[0]
        text = (choice.message.content or "").strip()
        truncated = getattr(choice, "finish_reason", None) == "length"
        usage = _normalize_usage("openai", getattr(resp, "usage", None))
    else:  # anthropic
        # Streamed, then collected into one final message. The Anthropic SDK
        # refuses NON-streaming requests whose max_tokens implies more than
        # ~10 minutes of generation (anything above ~21k tokens), raising
        # before the request is even sent — so with AI_MAX_TOKENS at 96k a
        # plain messages.create() can never succeed. get_final_message()
        # returns the same Message object (content, stop_reason, usage), so
        # everything below is unchanged.
        with client.messages.stream(
            model=model,
            max_tokens=AI_MAX_TOKENS,
            messages=[{"role": "user", "content": prompt}],
        ) as stream:
            resp = stream.get_final_message()
        text = "".join(block.text for block in resp.content
                       if getattr(block, "type", None) == "text").strip()
        truncated = getattr(resp, "stop_reason", None) == "max_tokens"
        usage = _normalize_usage("anthropic", getattr(resp, "usage", None))

    # Strip ```json ... ``` fences if the model added them despite instructions.
    if text.startswith("```"):
        text = re.sub(r"^```[a-zA-Z]*\s*", "", text)
        text = re.sub(r"\s*```$", "", text).strip()

    if truncated:
        # Don't raise; hand back the flag and the raw text so the caller can
        # salvage complete objects. Try a normal parse too, in case the JSON
        # happened to be complete despite the length finish_reason.
        try:
            return json.loads(text), usage, True, text
        except json.JSONDecodeError:
            return None, usage, True, text

    try:
        return json.loads(text), usage, False, text
    except json.JSONDecodeError as e:
        preview = text[:300].replace("\n", " ")
        raise ValueError(f"AI response was not valid JSON: {e}. "
                         f"First 300 chars: {preview!r}") from e


def _sum_usage(*totals_list: dict) -> dict:
    """Sum several usage accumulators into one."""
    combined = _new_usage_totals()
    for t in totals_list:
        for k in combined:
            combined[k] += t.get(k, 0)
    return combined


def _cost_for(totals: dict, model: str) -> float | None:
    """Estimated USD cost for one usage accumulator, or None if no pricing."""
    pricing = _model_pricing(model)
    if pricing is None:
        return None
    return (
        totals["input"] / 1_000_000 * pricing["input"]
        + totals["output"] / 1_000_000 * pricing["output"]
        + totals["cache_creation"] / 1_000_000 * pricing["cache_write"]
        + totals["cache_read"] / 1_000_000 * pricing["cache_read"]
    )


def _log_usage_block(label: str, totals: dict, model: str, logger: Logger):
    """Log one labelled usage block (a step or the total). Cache lines appear
    only when non-zero."""
    total_input = totals["input"] + totals["cache_creation"] + totals["cache_read"]
    logger.log(f"{label}:")
    logger.log(f"  API calls:      {totals['calls']}")
    logger.log(f"  Input tokens:   {total_input:,}")
    if totals["cache_read"] or totals["cache_creation"]:
        logger.log(f"    uncached:     {totals['input']:,}")
        logger.log(f"    cache writes: {totals['cache_creation']:,}")
        logger.log(f"    cache reads:  {totals['cache_read']:,}")
    logger.log(f"  Output tokens:  {totals['output']:,}")
    logger.log(f"  Total tokens:   {total_input + totals['output']:,}")
    cost = _cost_for(totals, model)
    if cost is None:
        logger.log("  Est. cost:      n/a (no pricing on file for this model)")
    else:
        logger.log(f"  Est. cost:      ${cost:,.4f} USD")


def _snapshot_records(records: list[dict]) -> list[dict]:
    """
    Copy the record list deeply enough to roll back a failed AI step.

    AI steps mutate records in two ways: the batched translation +
    definition step writes string fields in place, and cleaning rebuilds records
    on merge. A per-record dict copy therefore isn't enough on its own — the
    mutable internals ('_forms' set, '_example_pool' list) are copied too so a
    rollback cannot leave the restored list sharing state with the mutated one.
    The example dicts inside the pool are never modified in place, so they can
    be shared.
    """
    snap = []
    for r in records:
        c = dict(r)
        if "_forms" in c:
            c["_forms"] = set(c["_forms"])
        if "_example_pool" in c:
            c["_example_pool"] = list(c["_example_pool"])
        snap.append(c)
    return snap


def _log_usage_summary(clean_usage: dict, translate_usage: dict, model: str,
                       logger: Logger,
                       preflight_usage: dict | None = None) -> None:
    """Log per-step token usage (optional pre-flight language resolution,
    cleaning, translation + definitions) and the combined total. Steps with no
    calls are skipped. Nothing is logged if no AI calls were made at all."""
    pre = preflight_usage or _new_usage_totals()
    total = _sum_usage(pre, clean_usage, translate_usage)
    if total["calls"] == 0:
        return
    logger.log("--- AI token usage (this run) ---")
    logger.log(f"Model:            {model}")
    if pre["calls"]:
        _log_usage_block("Language resolution", pre, model, logger)
    if clean_usage["calls"]:
        _log_usage_block("Cleaning step", clean_usage, model, logger)
    if translate_usage["calls"]:
        _log_usage_block("Translate + define", translate_usage, model, logger)
    _log_usage_block("TOTAL", total, model, logger)
    logger.log("(estimate; verify against your provider billing)")
    logger.log("---------------------------------")


# ---------------------------------------------------------------------------
# AI STEP 1a: BUILD RECORDS FROM AGGREGATE OUTPUT
# ---------------------------------------------------------------------------

def _select_examples(pool: list[dict]) -> tuple[list[str], list[str]]:
    """
    Choose up to MAX_EXAMPLES examples from a term's pool, returning parallel
    (source_examples, target_examples) lists.

    Rules:
      - Prefer TMX-aligned examples (those with a non-empty target) over
        monolingual ones, so verified translations are surfaced first.
      - Otherwise keep first-seen order (representative of the corpus).
      - Fill up to MAX_EXAMPLES, TMX first, then monolingual to top up.
      - Each chosen example carries its own target verbatim (empty for
        monolingual), so conflicting TMX translations remain visible to the AI.
    """
    tmx = [e for e in pool if e.get("target")]
    mono = [e for e in pool if not e.get("target")]
    chosen = (tmx + mono)[:MAX_EXAMPLES]
    sources = [e["source"] for e in chosen]
    targets = [e.get("target", "") for e in chosen]
    return sources, targets


def _record_sort_key(rec: dict):
    """Global glossary ordering: most segments first, then raw frequency, then
    the term alphabetically so equal-ranked rows are stable and readable."""
    def _as_int(v):
        try:
            return int(v)
        except (TypeError, ValueError):
            return 0
    return (-_as_int(rec.get("seg_count")),
            -_as_int(rec.get("raw_freq")),
            str(rec.get("source_term", "")).lower())


def build_records(results_by_ngram: dict[int, list[dict]]) -> list[dict]:
    """
    Flatten the per-n-gram aggregate output into a single list of records in
    the glossary schema.

    The list is ordered globally by segment count (then raw frequency), NOT
    grouped by n-gram length, so the most important terms appear first however
    long they are. Each record carries its n-gram length in the 'ngram' column
    so the output can still be filtered by length in a spreadsheet. Ids are
    assigned after sorting, so they run in output order.

    Examples are selected (TMX-aligned preferred) into paired example_n /
    target_example_n columns; target_example_n is filled only for TMX-sourced
    examples. '_forms' (lowercased surface forms that collapsed into this
    record) and '_example_pool' (the full example dicts, for accurate
    re-selection on merge) are kept internally and stripped before output.
    """
    records: list[dict] = []
    for ngram_len in sorted(results_by_ngram.keys()):
        for entry in results_by_ngram[ngram_len]:
            pool = entry.get("examples", [])
            sources, targets = _select_examples(pool)
            rec = {
                "id": 0,     # assigned below, after the global sort
                "source_term": entry["term"],
                "target_term": "",
                "term_type": entry.get("term_type", ""),
                "ngram": ngram_len,
                "seg_count": entry["seg_count"],
                "raw_freq": entry["raw_freq"],
                "sources": entry.get("sources", ""),
                "definition": "",
                "notes": "",
                "_forms": {entry.get("_key", entry["term"].lower())},
                "_example_pool": pool,
            }
            for k in (1, 2, 3):
                rec[f"example_{k}"] = sources[k - 1] if k <= len(sources) else ""
                rec[f"target_example_{k}"] = targets[k - 1] if k <= len(targets) else ""
            records.append(rec)

    records.sort(key=_record_sort_key)
    for next_id, rec in enumerate(records):
        rec["id"] = next_id
    return records


# ---------------------------------------------------------------------------
# AI STEP 1b: CLEANING (whole-list, decisions-by-id)
# ---------------------------------------------------------------------------

def _records_for_cleaning_prompt(records: list[dict]) -> str:
    """Build the compact JSON array the cleaning prompt receives: id,
    source_term, counts, term_type, and the non-empty SOURCE examples.
    Target examples are deliberately omitted here: cleaning is a source-side
    judgement (dedup/normalise/drop) and targets would only waste tokens on
    the whole-list call. They are sent only in the translation step."""
    payload = []
    for r in records:
        examples = [r[f"example_{i}"] for i in (1, 2, 3) if r[f"example_{i}"]]
        payload.append({
            "id": r["id"],
            "source_term": r["source_term"],
            "seg_count": r["seg_count"],
            "raw_freq": r["raw_freq"],
            "term_type": r["term_type"],
            "examples": examples,
        })
    return json.dumps(payload, ensure_ascii=False, indent=2)


def _recompute_seg_count(forms: set[str], lowered_segments: list[str]) -> int:
    """Count distinct segments containing ANY of the given lowercased forms.
    This avoids the double-counting that plain summation would cause when two
    merged forms co-occur in the same segment."""
    count = 0
    for seg_lower in lowered_segments:
        if any(f in seg_lower for f in forms):
            count += 1
    return count


def apply_cleaning_decisions(records: list[dict], decisions: list[dict],
                             segments: list["Segment"],
                             logger: Logger) -> list[dict]:
    """
    Apply merge/drop/rename decisions to the record list, validating every
    decision against the set of ids actually sent. Any term not named in a
    valid decision passes through unchanged (implicit keep). seg_count is
    recomputed exactly for merges; raw_freq sums; sources union; examples pool.
    """
    by_id = {r["id"]: r for r in records}
    sent_ids = set(by_id)
    touched: set[int] = set()          # ids consumed by any valid decision
    dropped: set[int] = set()
    merged_away: set[int] = set()
    survivors: dict[int, dict] = {}     # id -> (possibly modified) record

    lowered_segments = [s.text.lower() for s in segments]

    def _conflict(ids: list[int]) -> bool:
        """True if any id is unknown or already consumed by another decision."""
        for i in ids:
            if i not in sent_ids:
                logger.log(f"[ai-clean] Ignoring decision referencing unknown "
                           f"id {i}.")
                return True
            if i in touched:
                logger.log(f"[ai-clean] Ignoring decision: id {i} already "
                           f"handled by an earlier decision.")
                return True
        return False

    n_merge = n_drop = n_rename = 0

    for dec in decisions:
        action = dec.get("action")

        if action == "drop":
            i = dec.get("id")
            if _conflict([i]):
                continue
            touched.add(i)
            dropped.add(i)
            n_drop += 1

        elif action == "rename":
            i = dec.get("id")
            new_term = (dec.get("canonical_term") or "").strip()
            if _conflict([i]):
                continue
            if not new_term:
                logger.log(f"[ai-clean] Ignoring rename of id {i}: empty "
                           f"canonical_term.")
                continue
            touched.add(i)
            rec = dict(by_id[i])
            _maybe_flag_canonical(new_term, rec, logger)
            rec["source_term"] = new_term
            rec["ngram"] = len(new_term.split())
            survivors[i] = rec
            n_rename += 1

        elif action == "merge":
            canon = dec.get("canonical_id")
            merge_ids = dec.get("merge_ids") or []
            new_term = (dec.get("canonical_term") or "").strip()
            if not isinstance(merge_ids, list) or not merge_ids:
                logger.log(f"[ai-clean] Ignoring merge with empty merge_ids "
                           f"(canonical_id={canon}).")
                continue
            if canon in merge_ids:
                logger.log(f"[ai-clean] Ignoring merge: canonical_id {canon} "
                           f"also appears in its own merge_ids.")
                continue
            group = [canon] + merge_ids
            if _conflict(group):
                continue
            for i in group:
                touched.add(i)
            merged_away.update(merge_ids)

            base = dict(by_id[canon])
            forms: set[str] = set()
            raw_freq = 0
            sources: set[str] = set()
            merged_pool: list[dict] = []
            for i in group:
                src = by_id[i]
                forms |= src["_forms"]
                raw_freq += src["raw_freq"]
                if src["sources"]:
                    sources.update(s.strip() for s in src["sources"].split(",")
                                   if s.strip())
                # Union the full example pools (not the pre-selected columns),
                # so a verified TMX example carried by a merged-away variant is
                # not lost. De-duplicate on (source, target).
                for ex in src.get("_example_pool", []):
                    if ex not in merged_pool:
                        merged_pool.append(ex)
            base["_forms"] = forms
            base["raw_freq"] = raw_freq
            # Recount distinct segments exactly against the corpus.
            base["seg_count"] = _recompute_seg_count(forms, lowered_segments)
            base["sources"] = ", ".join(sorted(sources))
            # Re-select over the combined pool: TMX-aligned examples are
            # preferred, so a verified translation surfaces regardless of which
            # merged variant it came from.
            base["_example_pool"] = merged_pool
            src_ex, tgt_ex = _select_examples(merged_pool)
            for k in (1, 2, 3):
                base[f"example_{k}"] = src_ex[k - 1] if k <= len(src_ex) else ""
                base[f"target_example_{k}"] = tgt_ex[k - 1] if k <= len(tgt_ex) else ""
            if new_term:
                _maybe_flag_canonical(new_term, base, logger)
                base["source_term"] = new_term
                base["ngram"] = len(new_term.split())
            survivors[canon] = base
            n_merge += 1

        else:
            logger.log(f"[ai-clean] Ignoring decision with unknown action "
                       f"{action!r}.")

    # Assemble the cleaned list: modified survivors + untouched implicit-keeps.
    cleaned: list[dict] = []
    for r in records:
        i = r["id"]
        if i in dropped or i in merged_away:
            continue
        cleaned.append(survivors.get(i, r))

    # Reconciliation accounting: every sent id must be classified exactly once.
    kept_untouched = sent_ids - touched
    accounted = kept_untouched | set(survivors) | merged_away | dropped
    orphans = sent_ids - accounted
    if orphans:
        logger.log(f"[ai-clean] WARNING: {len(orphans)} id(s) unaccounted for "
                   f"in reconciliation; kept unchanged: {sorted(orphans)}")

    logger.log(f"[ai-clean] Applied {n_merge} merge(s), {n_rename} rename(s), "
               f"{n_drop} drop(s). {len(records)} -> {len(cleaned)} term(s).")
    return cleaned


def _maybe_flag_canonical(new_term: str, rec: dict, logger: Logger):
    """Soft review flag: warn when a canonical term derives from none of the
    record's known surface forms (possible hallucination). Lenient: we still
    accept it, since legitimate normalisation (plural->singular) can produce a
    form matching no input exactly."""
    nt = new_term.lower()
    forms = rec.get("_forms", set())
    if not any(nt in f or f in nt for f in forms):
        logger.log(f"[ai-clean] Note: canonical term {new_term!r} does not "
                   f"derive from any extracted form {sorted(forms)}; kept for "
                   f"review.")


def ai_clean_terms(records: list[dict], segments: list["Segment"],
                   prompt_template: str, domain: str, context: str,
                   client, model: str, logger: Logger,
                   usage_totals: dict | None = None) -> list[dict]:
    """
    Whole-list cleaning step: build the prompt, call the model once,
    parse/salvage the decisions, and apply them. Returns the cleaned records.

    Failure semantics: an API error or a response with no 'decisions' array
    raises AIStepFailure, and the caller reverts to the pre-cleaning (NLP)
    state. A truncated response is not a failure: complete decisions are
    salvaged and applied.
    """
    label = "Cleaning"
    logger.log(f"[ai-clean] {label}: sending {len(records)} term(s) to "
               f"the model (single request)...")
    prompt = _fill_placeholders(prompt_template,
                                {"DOMAIN": domain, "CONTEXT": context})
    prompt += "\n\nHere is the term list:\n" + _records_for_cleaning_prompt(records)

    try:
        result, usage, truncated, raw_text = _call_ai_json(
            client, prompt, model, logger)
        if usage_totals is not None:
            _accumulate_usage(usage_totals, usage)
    except Exception as e:
        raise AIStepFailure(f"{label} failed "
                            f"({type(e).__name__}: {e})") from e

    if result is not None and isinstance(result.get("decisions"), list):
        decisions = result["decisions"]
        if truncated:
            logger.log_error(
                f"[ai-clean] {label}: response was truncated at the "
                f"{AI_MAX_TOKENS}-token limit, but its JSON happened to be "
                f"complete. Applied all {len(decisions)} decision(s).")
    elif truncated:
        salvaged = _salvage_json_array(raw_text, "decisions")
        if not salvaged:
            logger.log_error(
                f"[ai-clean] {label}: response was truncated at the "
                f"{AI_MAX_TOKENS}-token limit and no complete decisions could "
                f"be recovered. Keeping the list unchanged.")
            return records
        decisions = salvaged
        logger.log_error(
            f"[ai-clean] {label}: response was TRUNCATED at the "
            f"{AI_MAX_TOKENS}-token limit. Recovered and applied "
            f"{len(decisions)} complete decision(s) from the partial "
            f"response; remaining terms were left uncleaned. To clean the "
            f"whole list, use a smaller run.")
    else:
        raise AIStepFailure(f"{label}: response had no 'decisions' array")

    return apply_cleaning_decisions(records, decisions, segments, logger)


# ---------------------------------------------------------------------------
# AI STEP 2: TRANSLATE + DEFINE (batched, one request per batch)
# ---------------------------------------------------------------------------

def _records_for_translation_prompt(batch: list[dict]) -> str:
    """Compact JSON for a monolingual translate + define batch: id,
    source_term, and non-empty source examples."""
    payload = []
    for r in batch:
        examples = [r[f"example_{i}"] for i in (1, 2, 3) if r[f"example_{i}"]]
        payload.append({"id": r["id"], "source_term": r["source_term"],
                        "examples": examples})
    return json.dumps(payload, ensure_ascii=False, indent=2)


def _records_for_tmx_prompt(batch: list[dict]) -> str:
    """Compact JSON for a TMX extraction batch: id, source_term, and paired
    (source, target) example segments so the model can read the verified
    translation. Only pairs with a non-empty target are sent as 'aligned';
    any source-only examples are included separately as plain context."""
    payload = []
    for r in batch:
        aligned = []
        context_only = []
        for i in (1, 2, 3):
            src = r[f"example_{i}"]
            tgt = r[f"target_example_{i}"]
            if src and tgt:
                aligned.append({"source": src, "target": tgt})
            elif src:
                context_only.append(src)
        entry = {
            "id": r["id"],
            "source_term": r["source_term"],
            "aligned_examples": aligned,
        }
        if context_only:
            entry["source_only_examples"] = context_only
        payload.append(entry)
    return json.dumps(payload, ensure_ascii=False, indent=2)


def _apply_translation_results(result: dict, by_id: dict, batch_label: str,
                               logger: Logger) -> int:
    """Join a translate + define (or TMX extraction + define) response back
    to records by id. Fills target_term, definition and (if present) notes.
    Returns count filled."""
    translations = result.get("translations")
    if not isinstance(translations, list):
        raise AIStepFailure(f"{batch_label}: response had no 'translations' "
                            f"array")
    filled = 0
    for item in translations:
        i = item.get("id")
        if i not in by_id:
            logger.log(f"[ai-translate] {batch_label}: ignoring unknown id {i}.")
            continue
        rec = by_id[i]
        rec["target_term"] = (item.get("target_term") or "").strip()
        rec["definition"] = (item.get("definition") or "").strip()
        note = (item.get("notes") or "").strip()
        if note:
            rec["notes"] = note
        # TMX path: bold the term's surface form(s) in the target examples.
        forms = item.get("target_surface_forms")
        if isinstance(forms, list) and forms:
            for k in (1, 2, 3):
                tgt = rec.get(f"target_example_{k}", "")
                if tgt:
                    rec[f"target_example_{k}"] = _bold_forms_in_text(tgt, forms)
        filled += 1
    return filled


def _run_translation_pass(records: list[dict], base_prompt: str,
                          payload_builder, pass_label: str, client, model: str,
                          logger: Logger, usage_totals: dict | None) -> int:
    """Batch a set of records through one prompt and apply the results.
    Returns the number of terms filled."""
    total = len(records)
    if total == 0:
        return 0
    by_id = {r["id"]: r for r in records}
    n_batches = (total + TRANSLATE_BATCH_SIZE - 1) // TRANSLATE_BATCH_SIZE
    logger.log(f"[ai-translate] {pass_label}: {total} term(s) in "
               f"{n_batches} batch(es).")
    filled = 0
    for b in range(n_batches):
        batch = records[b * TRANSLATE_BATCH_SIZE:(b + 1) * TRANSLATE_BATCH_SIZE]
        prompt = (base_prompt + "\n\nHere are the terms:\n"
                  + payload_builder(batch))
        label = f"{pass_label} batch {b + 1}/{n_batches}"
        logger.log(f"[ai-translate] {label} ({len(batch)} term(s))...")
        try:
            result, usage, truncated, raw_text = _call_ai_json(
                client, prompt, model, logger)
            if usage_totals is not None:
                _accumulate_usage(usage_totals, usage)
        except Exception as e:
            raise AIStepFailure(f"{label} failed "
                                f"({type(e).__name__}: {e})") from e
        if result is None and truncated:
            # Salvage whatever complete translation objects we can; the rest of
            # this batch's terms simply stay untranslated.
            salvaged = _salvage_json_array(raw_text, "translations")
            if salvaged:
                result = {"translations": salvaged}
                logger.log_error(
                    f"[ai-translate] {label} was TRUNCATED at the "
                    f"{AI_MAX_TOKENS}-token limit; recovered "
                    f"{len(salvaged)} of {len(batch)} term(s). The rest were "
                    f"left untranslated and undefined.")
            else:
                logger.log_error(
                    f"[ai-translate] {label} was TRUNCATED at the "
                    f"{AI_MAX_TOKENS}-token limit and nothing could be "
                    f"recovered; those terms were left untranslated and "
                    f"undefined.")
                continue
        elif truncated:
            logger.log_error(
                f"[ai-translate] {label} hit the {AI_MAX_TOKENS}-token limit "
                f"but returned complete JSON.")
        filled += _apply_translation_results(result, by_id, label, logger)
    return filled


def ai_translate_terms(records: list[dict], mono_prompt: str, tmx_prompt: str,
                       domain: str, context: str, target_language: str,
                       definition_language: str, client, model: str,
                       logger: Logger,
                       usage_totals: dict | None = None) -> list[dict]:
    """
    Translate + define: fill target_term, definition and notes for every
    record, one request per batch returning both the translation and the
    definition.

    Records are partitioned by whether they carry any TMX-aligned example
    (a non-empty target_example): those go through the TMX 'extraction' prompt,
    which asks the model to read the verified human translation from the target
    segment; the rest go through the standard monolingual translation prompt.
    Both partitions are batched (TRANSLATE_BATCH_SIZE), resilient per-batch, and
    write to the same fields. Usage is accumulated across both.
    """
    total = len(records)
    if total == 0:
        return records

    placeholders = {
        "DOMAIN": domain,
        "CONTEXT": context,
        "TARGET_LANGUAGE": target_language,
        "DEFINITION_LANGUAGE": definition_language,
    }
    mono_base = _fill_placeholders(mono_prompt, placeholders)
    tmx_base = _fill_placeholders(tmx_prompt, placeholders)

    def has_target(r: dict) -> bool:
        return any(r[f"target_example_{i}"] for i in (1, 2, 3))

    tmx_records = [r for r in records if has_target(r)]
    mono_records = [r for r in records if not has_target(r)]

    logger.log(f"[ai-translate] Translating {total} term(s) into "
               f"{target_language}, with definitions in "
               f"{definition_language}: "
               f"{len(tmx_records)} via TMX extraction, {len(mono_records)} via "
               f"standard translation.")

    filled = 0
    filled += _run_translation_pass(
        tmx_records, tmx_base, _records_for_tmx_prompt, "TMX extraction",
        client, model, logger, usage_totals)
    filled += _run_translation_pass(
        mono_records, mono_base, _records_for_translation_prompt,
        "Translation", client, model, logger, usage_totals)

    missing = total - filled
    if missing:
        logger.log(f"[ai-translate] {missing} term(s) left without a "
                   f"translation/definition (failed/omitted).")
    logger.log(f"[ai-translate] Done: {filled}/{total} term(s) translated "
               f"and defined.")
    return records


# ---------------------------------------------------------------------------
# STEP 4: AGGREGATION
# ---------------------------------------------------------------------------

def aggregate(segments: list[Segment], spacy_candidates: dict[str, dict],
              selected_ngrams: set[int],
              term_limit: int | None, min_count: dict[int, int],
              logger: Logger, limit_mode: str = "per_ngram",
              unspaced: bool = False) -> dict[int, list[dict]]:
    """
    For each candidate term (keyed by cleaned lowercased form) compute:
      - term      : the most frequent surface spelling (best form)
      - term_type : Noun / Adjective / Verb / Other
      - seg_count : number of distinct segments containing it (any source)
      - raw_freq  : total occurrences
      - sources   : set of source types it appears in
      - examples  : up to MAX_EXAMPLES sentences, term in **bold**
    Filter by min_count, group by n-gram length, sort by seg_count,
    optionally cap to term_limit per n-gram.

    unspaced=True (Chinese, Japanese): terms are grouped by CHARACTER length
    instead of word count, no length is filtered out (no n-gram selection, no
    5-word ceiling), a single min. count (the 1-gram field) applies to all
    lengths, and only a total cap with global ranking is used.
    """
    if unspaced:
        limit_mode = "total"
    candidates = set(spacy_candidates.keys())
    logger.log(f"[aggregate] {len(candidates)} candidate(s) to aggregate.")

    # Pre-lowercase segments once, keeping source tags
    lowered = [(s.text.lower(), s.source) for s in segments]

    results_by_ngram: dict[int, list[dict]] = defaultdict(list)

    skipped: Counter = Counter()
    offered_max = max(NGRAM_CHOICES)
    single_min = min_count.get(1, 1)
    for key in candidates:
        if unspaced:
            ngram_len = _term_char_length(key)
        else:
            ngram_len = len(key.split())
        if not unspaced and ngram_len not in selected_ngrams:
            # Distinguish a length the user merely left unticked (recoverable
            # by ticking it) from one the tool cannot offer at all.
            if ngram_len > offered_max:
                skipped[f"longer than {offered_max} words"] += 1
            else:
                skipped[f"{ngram_len}-gram not selected"] += 1
            continue

        seg_count = 0
        raw_freq = 0
        sources: set[str] = set()
        for seg_lower, src in lowered:
            c = seg_lower.count(key)
            if c > 0:
                seg_count += 1
                raw_freq += c
                sources.add(src)

        threshold = single_min if unspaced else min_count.get(ngram_len, 5)
        if seg_count < threshold:
            skipped["below min. count"] += 1
            continue

        meta = spacy_candidates[key]
        best_form = meta["forms"].most_common(1)[0][0] if meta["forms"] else key
        top_pos = meta["pos_counts"].most_common(1)[0][0] if meta["pos_counts"] else None

        results_by_ngram[ngram_len].append({
            "term": best_form,
            "term_type": _term_type_label(top_pos),
            "seg_count": seg_count,
            "raw_freq": raw_freq,
            "sources": ", ".join(sorted(sources)),
            "examples": meta["examples"],   # full pool; selected in build_records
            "_key": key,   # lowercased canonical key, for AI merge recomputation
        })

    total_skipped = sum(skipped.values())
    if total_skipped:
        detail = ", ".join(f"{r}: {n}" for r, n in skipped.most_common())
        logger.log(f"[aggregate] Skipped {total_skipped} of "
                   f"{len(candidates)} candidate(s) — {detail}.")

    # Rank within each n-gram length first (shared ranking key throughout).
    _rank = lambda x: (-x["seg_count"], -x["raw_freq"])
    for ngram_len in sorted(results_by_ngram):
        results_by_ngram[ngram_len].sort(key=_rank)
        if not unspaced:
            logger.log(f"[aggregate] n={ngram_len}: "
                       f"{len(results_by_ngram[ngram_len])} term(s) above "
                       f"threshold.")
    if unspaced:
        logger.log(f"[aggregate] "
                   f"{sum(len(v) for v in results_by_ngram.values())} term(s) "
                   f"above threshold (all character lengths).")

    if term_limit is None:
        return results_by_ngram

    if limit_mode == "per_ngram":
        for ngram_len in sorted(results_by_ngram):
            total = len(results_by_ngram[ngram_len])
            if total > term_limit:
                results_by_ngram[ngram_len] = \
                    results_by_ngram[ngram_len][:term_limit]
                logger.log(f"[aggregate] n={ngram_len}: capped to top "
                           f"{term_limit} of {total}.")
        return results_by_ngram

    # --- total cap ---
    grand_total = sum(len(v) for v in results_by_ngram.values())
    if grand_total <= term_limit:
        logger.log(f"[aggregate] {grand_total} term(s) total, under the cap "
                   f"of {term_limit}; nothing dropped.")
        return results_by_ngram

    # Global sort-and-truncate: pool everything, rank together, keep top N.
    pooled = [(n, entry) for n, entries in results_by_ngram.items()
              for entry in entries]
    pooled.sort(key=lambda pair: _rank(pair[1]))
    kept = pooled[:term_limit]
    new_results = {n: [] for n in results_by_ngram}
    for n, entry in kept:
        new_results[n].append(entry)
    for n in sorted(new_results):
        new_results[n].sort(key=_rank)
        if not unspaced:
            logger.log(f"[aggregate] n={n}: kept {len(new_results[n])} of "
                       f"{len(results_by_ngram[n])} (global ranking).")
    logger.log(f"[aggregate] Total capped to {len(kept)} term(s) from "
               f"{grand_total}.")
    return new_results


# ---------------------------------------------------------------------------
# STEP 5: OUTPUT (JSON + mirrored XLSX)
# ---------------------------------------------------------------------------

HEADER_FILL = PatternFill("solid", fgColor="2E4057")
HEADER_FONT = Font(bold=True, color="FFFFFF")
ALT_FILL = PatternFill("solid", fgColor="F2F2F2")

_BOLD_SPLIT = re.compile(r"(\*\*.*?\*\*)")

def _markdown_to_rich(text: str):
    """
    Convert a string with **bold** markers into an openpyxl CellRichText
    with real bold runs. If there are no markers, return the plain string.
    """
    if "**" not in text:
        return text
    parts = [p for p in _BOLD_SPLIT.split(text) if p]
    runs = []
    for part in parts:
        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
            runs.append(TextBlock(InlineFont(b=True), part[2:-2]))
        else:
            runs.append(part)
    return CellRichText(*runs)


def records_to_output(records: list[dict], by_chars: bool = False) -> list[dict]:
    """Project internal records onto the public schema (GLOSSARY_COLUMNS),
    dropping internal helper keys ('id', '_forms', '_example_pool').

    Sorted here as well as at build time, because the AI cleaning step can
    merge terms and change their seg_count — so the final file is always
    ordered by importance regardless of what happened upstream."""
    out = []
    for r in sorted(records, key=_record_sort_key):
        row = {}
        for col in GLOSSARY_COLUMNS:
            if col == "ngram" and by_chars:
                # Recomputed from the final term, so AI renames/merges are
                # always reflected.
                row[LENGTH_COLUMN_CHARS] = _term_char_length(
                    str(r.get("source_term", "")))
            else:
                row[col] = r.get(col, "")
        out.append(row)
    return out


def write_json(records: list[dict], output_path: str, logger: Logger,
               by_chars: bool = False):
    """Write the glossary as a JSON array in the canonical schema."""
    data = records_to_output(records, by_chars)
    with open(output_path, "w", encoding="utf-8") as fh:
        json.dump(data, fh, ensure_ascii=False, indent=2)
    logger.log(f"[json] Saved {len(data)} term(s) to '{Path(output_path).name}'.")


def write_xlsx(records: list[dict], output_path: str, logger: Logger,
               by_chars: bool = False):
    """Write a flat 'Glossary' sheet mirroring the JSON schema, column-for-
    column. Example columns render **bold** markers as real bold runs."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Glossary"

    data = records_to_output(records, by_chars)

    if not data:
        ws["A1"] = ("No terms passed the filters. Try lowering the "
                    "'Min. count per n-gram' values or check the source language.")
        ws["A1"].font = Font(bold=True, color="B00020")
        ws.column_dimensions["A"].width = 90
        wb.save(output_path)
        logger.log(f"[xlsx] No terms — wrote an empty report to "
                   f"'{Path(output_path).name}'.")
        return

    headers = _output_columns(by_chars)
    example_cols = {"example_1", "example_2", "example_3",
                    "target_example_1", "target_example_2", "target_example_3"}

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    for row_idx, entry in enumerate(data, start=2):
        fill = ALT_FILL if row_idx % 2 == 0 else None
        for col_idx, header in enumerate(headers, start=1):
            value = entry.get(header, "")
            if header in example_cols and isinstance(value, str):
                value = _markdown_to_rich(value)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if fill:
                cell.fill = fill
            if header in example_cols or header in ("definition", "notes"):
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    widths = {
        "source_term": 24, "target_term": 24, "term_type": 11, "ngram": 8,
        LENGTH_COLUMN_CHARS: 11,
        "seg_count": 10, "raw_freq": 10, "sources": 16, "definition": 45,
        "example_1": 50, "target_example_1": 50,
        "example_2": 50, "target_example_2": 50,
        "example_3": 50, "target_example_3": 50,
        "notes": 30,
    }
    for col_idx, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = widths.get(header, 20)
    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"

    wb.save(output_path)
    logger.log(f"[xlsx] Saved {len(data)} term(s) to '{Path(output_path).name}'.")


# ---------------------------------------------------------------------------
# PIPELINE DRIVER
# ---------------------------------------------------------------------------

def run_pipeline(folder: str, language: str, tmx_paths: list[Path],
                 pdf_paths: list[Path], docx_paths: list[Path],
                 selected_ngrams: set[int], term_limit: int | None,
                 custom_stopwords: set, min_count: dict[int, int],
                 min_segment_tokens: int, logger: Logger,
                 *, ai_clean: bool = False, ai_translate: bool = False,
                 api_key: str = "", ai_model: str = "",
                 domain: str = "", context: str = "",
                 target_language: str = "", definition_language: str = "",
                 tmx_target_codes: dict | None = None,
                 preflight_usage: dict | None = None,
                 limit_mode: str = "per_ngram", junk_filter: bool = True):
    logger.log(f"=== Terminology Extractor v{VERSION} ===")
    start_time = datetime.now()
    stamp = start_time.strftime("%Y-%m-%d_%H-%M-%S")
    unspaced = language in UNSPACED_LANGUAGES
    if unspaced:
        # No word count to cap by: total cap with global ranking only.
        limit_mode = "total"
    by_chars = unspaced
    folder_path = Path(folder)
    out_dir = folder_path
    output_base = f"{folder_path.name}_term_extraction_{stamp}"
    json_path = str(out_dir / f"{output_base}.json")
    xlsx_path = str(out_dir / f"{output_base}.xlsx")

    # Log the settings used for this run.
    steps = []
    if ai_clean:
        steps.append("clean")
    if ai_translate:
        steps.append("translate + define")
    ai_str = f"{' + '.join(steps)} ({ai_model})" if steps else "off (NLP only)"
    logger.log(f"Input folder:        {folder}")
    logger.log(f"Output (JSON):       {json_path}")
    logger.log(f"Output (XLSX):       {xlsx_path}")
    logger.log("--- Settings ---")
    ngrams_str = ", ".join(f"{n}" for n in sorted(selected_ngrams))
    mincount_str = ", ".join(f"{n}g={min_count.get(n)}"
                             for n in sorted(selected_ngrams))
    if term_limit is None:
        limit_str = "no limit"
    elif limit_mode == "total":
        limit_str = f"{term_limit} total (global ranking)"
    else:
        limit_str = f"{term_limit} per n-gram"
    stopwords_str = (", ".join(sorted(custom_stopwords))
                     if custom_stopwords else "none")
    if unspaced:
        ngrams_str = ("not applied (all lengths kept; length measured "
                      "in characters)")
        mincount_str = f"{min_count.get(1, 1)} (all lengths)"
    logger.log(f"Source language:     {language}")
    logger.log(f"N-grams:             {ngrams_str}")
    if unspaced:
        logger.log("Min. segment length: not applied (unspaced language)")
    else:
        logger.log(f"Min. segment length: {min_segment_tokens} token(s)")
    logger.log(f"Min. count:          {mincount_str}")
    logger.log(f"Max terms:           {limit_str}")
    logger.log(f"Custom stopwords:    {stopwords_str}")
    logger.log(f"Junk filter:         "
               f"{'on' if junk_filter else 'off'}")
    logger.log(f"AI enhancement:      {ai_str}")
    if ai_translate:
        logger.log(f"Target language:     {target_language}")
        logger.log(f"Definition language: {definition_language}")
    logger.log("----------------\n")

    if not (tmx_paths or pdf_paths or docx_paths):
        logger.log("No files to process. Nothing to do.")
        return

    nlp = load_spacy_model(language, logger)

    # Collect segments from every source type
    segments = []
    segments += parse_tmx_files(tmx_paths, logger, tmx_target_codes,
                                will_translate=ai_translate)
    segments += parse_pdf_files(pdf_paths, nlp, logger)
    segments += parse_docx_files(docx_paths, nlp, logger)
    logger.log(f"[collect] {len(segments)} total segment(s) from all "
               f"sources.\n")

    segments = clean_segments(segments, logger)

    spacy_candidates = extract_candidates_spacy(
        segments, nlp, selected_ngrams, custom_stopwords,
        min_segment_tokens, logger, junk_filter=junk_filter)
    results = aggregate(segments, spacy_candidates,
                        selected_ngrams, term_limit, min_count, logger,
                        limit_mode=limit_mode, unspaced=unspaced)

    # Flatten into the flat glossary schema with stable ids.
    records = build_records(results)
    logger.log(f"[records] {len(records)} candidate term(s) in the "
               f"glossary.\n")

    # --- Optional AI enhancement ---
    # Failure policy: an AI step that cannot complete (API error, unparseable
    # JSON, missing array) halts the AI chain. The failed step's partial work
    # is DISCARDED and the output is written from the last fully completed
    # step. Truncation is not a failure — it salvages and continues.
    clean_usage = _new_usage_totals()
    translate_usage = _new_usage_totals()
    pre_usage = preflight_usage or _new_usage_totals()
    if ai_clean:
        last_good_stage = "NLP extraction"
        try:
            prompts = load_prompts(logger)
            client = make_ai_client(api_key, _model_provider(ai_model))

            # --- Cleaning ---
            snapshot = _snapshot_records(records)
            try:
                records = ai_clean_terms(records, segments, prompts["cleaning"],
                                         domain, context, client, ai_model,
                                         logger, usage_totals=clean_usage)
                last_good_stage = "AI cleaning"
            except AIStepFailure as e:
                records = snapshot
                raise AIStepFailure(f"cleaning step: {e}") from e

            # --- Translate + define ---
            if ai_translate:
                snapshot = _snapshot_records(records)
                try:
                    records = ai_translate_terms(
                        records, prompts["translation"],
                        prompts["tmx_translation"], domain, context,
                        target_language, definition_language, client,
                        ai_model, logger, usage_totals=translate_usage)
                    last_good_stage = "AI translation + definitions"
                except AIStepFailure as e:
                    records = snapshot
                    raise AIStepFailure(f"translate + define step: {e}") from e

        except AIStepFailure as e:
            logger.log_error(
                f"[ai] STOPPED: {e}. No further AI steps were run and the "
                f"failed step's partial results were discarded. The output "
                f"below reflects the last completed stage: {last_good_stage}.")
        except Exception as e:
            logger.log_error(
                f"[ai] STOPPED: AI enhancement could not run "
                f"({type(e).__name__}: {e}). The output below reflects the "
                f"last completed stage: {last_good_stage}.")
        _log_usage_summary(clean_usage, translate_usage, ai_model, logger,
                           preflight_usage=pre_usage)
    elif pre_usage["calls"]:
        # Pre-flight used the API even though cleaning is off; still report it.
        _log_usage_summary(_new_usage_totals(), _new_usage_totals(),
                           ai_model, logger, preflight_usage=pre_usage)

    # --- Output: always both JSON and mirrored XLSX ---
    write_json(records, json_path, logger, by_chars)
    write_xlsx(records, xlsx_path, logger, by_chars)

    elapsed = datetime.now() - start_time
    total_seconds = int(elapsed.total_seconds())
    minutes, seconds = divmod(total_seconds, 60)
    duration = f"{minutes}m {seconds}s" if minutes else f"{seconds}s"
    logger.log(f"\n=== Done in {duration} ===")


# ---------------------------------------------------------------------------
# GUI
# ---------------------------------------------------------------------------

class Tooltip:
    """Simple hover tooltip for a Tkinter widget."""
    def __init__(self, widget, text: str, wraplength: int = 320):
        self.widget = widget
        self.text = text
        self.wraplength = wraplength
        self.tipwindow = None
        widget.bind("<Enter>", self._show)
        widget.bind("<Leave>", self._hide)

    def _show(self, event=None):
        if self.tipwindow or not self.text:
            return
        self.tipwindow = tw = tk.Toplevel(self.widget)
        tw.wm_overrideredirect(True)   # no window border
        label = tk.Label(tw, text=self.text, justify="left",
                         background="#ffffe0", relief="solid", borderwidth=1,
                         wraplength=self.wraplength, padx=8, pady=6)
        label.pack()

        # Default position: just below and slightly right of the widget.
        x = self.widget.winfo_rootx() + 20
        y = self.widget.winfo_rooty() + self.widget.winfo_height() + 6

        # Keep the tip fully on screen: if placing it at x would push its right
        # edge past the screen, shift it left so it fits (with an 8px margin).
        tw.update_idletasks()   # so winfo_width/reqwidth are accurate
        tip_w = tw.winfo_reqwidth()
        screen_w = self.widget.winfo_screenwidth()
        if x + tip_w + 8 > screen_w:
            x = max(8, screen_w - tip_w - 8)
        # Guard the bottom edge too, flipping above the widget if needed.
        tip_h = tw.winfo_reqheight()
        screen_h = self.widget.winfo_screenheight()
        if y + tip_h + 8 > screen_h:
            y = max(8, self.widget.winfo_rooty() - tip_h - 6)

        tw.wm_geometry(f"+{x}+{y}")

    def _hide(self, event=None):
        if self.tipwindow:
            self.tipwindow.destroy()
            self.tipwindow = None


class App:
    LIMIT_PLACEHOLDER = "No limit"

    def __init__(self, root):
        self.root = root
        self.root.title(f"Terminology Extractor v{VERSION} (TMX / PDF / DOCX)")
        self.root.geometry("820x900")

        self.selected_path = tk.StringVar(value="")
        self.recurse_var = tk.BooleanVar(value=True)
        self.language_var = tk.StringVar(value=DEFAULT_LANGUAGE)
        self.ngram_vars = {n: tk.BooleanVar(value=(n != 5))
                           for n in NGRAM_CHOICES}
        # Term cap: "per_ngram" (cap each length separately) or "total"
        # (cap the whole glossary). Total is the more natural control.
        self.limit_mode_var = tk.StringVar(value="total")
        self.junk_filter_var = tk.BooleanVar(value=True)
        # --- AI enhancement state ---
        self.ai_clean_var = tk.BooleanVar(value=True)
        self.ai_translate_var = tk.BooleanVar(value=True)
        self.ai_model_var = tk.StringVar(value=DEFAULT_AI_MODEL_LABEL)
        # Seed the key from the environment variable matching the default
        # model's provider (ANTHROPIC_API_KEY for Claude, OPENAI_API_KEY for
        # GPT). Re-seeded when the model changes, unless the user typed a key.
        self.api_key_var = tk.StringVar(
            value=self._env_key_for_label(DEFAULT_AI_MODEL_LABEL))
        self._api_key_user_edited = False
        self.target_language_var = tk.StringVar(value="")
        self.definition_language_var = tk.StringVar(value="English")
        # Min-count placeholders shown in grey in the GUI
        # Min. count defaults to 1 (i.e. no frequency filtering): meaningful
        # thresholds depend entirely on corpus size, so we don't guess.
        self.min_count_defaults = {1: 1, 2: 1, 3: 1, 4: 1, 5: 1}
        self.min_count_entries: dict[int, ttk.Entry] = {}
        self.min_segment_tokens_default = MIN_SEGMENT_TOKENS
        self.log_queue: queue.Queue = queue.Queue()
        self.worker = None
        self._last_line_updatable = False

        self._build_ui()
        self._poll_log_queue()

    def _build_ui(self):
        pad = {"padx": 10, "pady": 6}

        # Bold titles for the main sections.
        style = ttk.Style()
        _base = tkfont.nametofont("TkDefaultFont")
        style.configure("Section.TLabelframe.Label",
                        font=(_base.actual("family"),
                              _base.actual("size"), "bold"))

        # Row 1: folder path + browse + subfolder toggle
        path_frame = ttk.LabelFrame(self.root, text="Input",
                                    style="Section.TLabelframe")
        path_frame.pack(fill="x", **pad)

        inner = ttk.Frame(path_frame)
        inner.pack(fill="x", padx=8, pady=6)
        self.path_entry = ttk.Entry(inner, textvariable=self.selected_path)
        self.path_entry.pack(side="left", fill="x", expand=True, padx=(0, 8))
        self.browse_btn = ttk.Button(inner, text="Browse…", command=self._on_browse)
        self.browse_btn.pack(side="left")
        self.recurse_check = ttk.Checkbutton(
            inner, text="Scan subfolders", variable=self.recurse_var)
        self.recurse_check.pack(side="left", padx=(10, 0))

        # Row 2: options (language + n-grams + limit + stopwords + min counts)
        opts_frame = ttk.LabelFrame(self.root, text="NLP extraction options",
                                    style="Section.TLabelframe")
        opts_frame.pack(fill="x", **pad)

        # -- language + n-grams line --
        line0 = ttk.Frame(opts_frame)
        line0.pack(fill="x", padx=8, pady=(6, 2))
        ttk.Label(line0, text="Source language:").pack(side="left", padx=(0, 6))
        lang_values = [DEFAULT_LANGUAGE] + [
            lang for lang in LANGUAGE_MODELS if lang != DEFAULT_LANGUAGE]
        self.language_combo = ttk.Combobox(
            line0, textvariable=self.language_var,
            values=lang_values, state="readonly", width=20)
        self.language_combo.pack(side="left")
        self.language_combo.bind("<<ComboboxSelected>>",
                                 lambda e: self._refresh_language_states())

        lang_info = ttk.Label(line0, text="ⓘ", foreground="#2E4057",
                              font=("TkDefaultFont", 11, "bold"),
                              cursor="question_arrow")
        lang_info.pack(side="left", padx=(8, 0))
        Tooltip(lang_info, "Language of the SOURCE text to extract terms from. "
                           "The list is limited to the languages with an "
                           "available spaCy model, so a language missing here "
                           "cannot be used as the extraction source. This does "
                           "not restrict the target language for translations, "
                           "which is free text.\n\n"
                           "Chinese and Japanese are written without spaces, "
                           "so word-based options don't apply: all term "
                           "lengths are kept, the length column is written "
                           "as 'characters', Min. segment length is off, one "
                           "Min. count applies to all lengths, and Max terms "
                           "is a total cap only.")

        ttk.Label(line0, text="N-grams:").pack(side="left", padx=(20, 6))
        self.ngram_checks = {}
        for n in NGRAM_CHOICES:
            chk = ttk.Checkbutton(line0, text=f"{n}-gram{'s' if n > 1 else ''}",
                                  variable=self.ngram_vars[n])
            chk.pack(side="left", padx=3)
            self.ngram_checks[n] = chk

        # -- max-terms line (per n-gram vs total) --
        line_cap = ttk.Frame(opts_frame)
        line_cap.pack(fill="x", padx=8, pady=(2, 2))
        ttk.Label(line_cap, text="Max terms:").pack(side="left", padx=(0, 6))
        self.per_ngram_radio = ttk.Radiobutton(
            line_cap, text="per n-gram",
            variable=self.limit_mode_var, value="per_ngram")
        self.per_ngram_radio.pack(side="left")
        ttk.Radiobutton(line_cap, text="total",
                        variable=self.limit_mode_var, value="total").pack(
            side="left", padx=(8, 0))
        self.limit_entry = ttk.Entry(line_cap, width=10)
        self.limit_entry.pack(side="left", padx=(10, 0))
        self.limit_entry.insert(0, self.LIMIT_PLACEHOLDER)
        self.limit_entry.config(foreground="grey")
        self.limit_entry.bind("<FocusIn>", self._on_limit_focus_in)
        self.limit_entry.bind("<FocusOut>", self._on_limit_focus_out)
        self.limit_entry.bind("<KeyRelease>", self._on_limit_key)

        limit_info = ttk.Label(line_cap, text="ⓘ", foreground="#2E4057",
                               font=("TkDefaultFont", 11, "bold"),
                               cursor="question_arrow")
        limit_info.pack(side="left", padx=(8, 0))
        Tooltip(limit_info,
                "Caps how many terms are kept, keeping the highest-ranked "
                "ones (by segment count, then raw frequency).\n\n"
                "'per n-gram' applies the cap separately to each n-gram "
                "length. 'total' caps the whole glossary: all terms are "
                "pooled, ranked together, and the lowest-ranked are dropped. "
                "Note that a total cap tends to favour 1-grams, which "
                "naturally have higher counts — untick 1-gram, or use a "
                "per n-gram cap, for a more balanced mix.\n\n"
                "For Chinese and Japanese only 'total' is available, with "
                "global ranking: there is no word count to cap by.\n\n"
                "Leave blank for no limit.")

        # -- Advanced sub-section: options that silently discard data if set
        #    wrong, so they're grouped apart from the everyday controls.
        adv_frame = ttk.LabelFrame(opts_frame, text="Advanced")
        adv_frame.pack(fill="x", padx=8, pady=(6, 8))

        # -- junk filter line --
        adv_junk = ttk.Frame(adv_frame)
        adv_junk.pack(fill="x", padx=8, pady=(6, 6))
        self.junk_filter_check = ttk.Checkbutton(
            adv_junk, text="Filter out extraction junk",
            variable=self.junk_filter_var)
        self.junk_filter_check.pack(side="left")
        junk_info = ttk.Label(adv_junk, text="ⓘ", foreground="#2E4057",
                              font=("TkDefaultFont", 11, "bold"),
                              cursor="question_arrow")
        junk_info.pack(side="left", padx=(8, 0))
        Tooltip(junk_info,
                "Discards candidates that are clearly not terms, before they "
                "enter the pipeline: document and part codes (LBL-02011, "
                "GU14), phone numbers and long digit strings, bare "
                "measurements (1 cm, less than 8 cm, up to 10 ml), standalone "
                "units, and truncated fragments with unbalanced brackets or "
                "quotes ('(inches', 'code \"US').\n\n"
                "The rules are deliberately cautious — terms that merely "
                "contain a number, such as '10ml syringe' or "
                "'120º anteversion', are kept. Turn this off to see the raw "
                "candidate list. The run log reports what was removed and "
                "why.\n\n"
                "The rules work in any script. For Chinese and Japanese, "
                "one- and two-character terms are kept, and characters next "
                "to digits ('3号機', '12月') are not treated as codes; codes "
                "written in Latin letters are still caught.")

        # -- min. segment length (same row as the junk filter) --
        adv_line0 = adv_junk
        ttk.Label(adv_line0, text="Min. segment length:").pack(
            side="left", padx=(20, 6))
        self.mintok_entry = ttk.Entry(adv_line0, width=4)
        self.mintok_entry.pack(side="left")
        self.mintok_entry.insert(0, str(self.min_segment_tokens_default))
        self.mintok_entry.config(foreground="grey")
        self._mintok_is_placeholder = True
        self.mintok_entry.bind("<FocusIn>", self._on_mintok_focus_in)
        self.mintok_entry.bind("<FocusOut>", self._on_mintok_focus_out)
        self.mintok_entry.bind("<KeyRelease>", self._on_mintok_key)

        mintok_info = ttk.Label(adv_line0, text="ⓘ", foreground="#2E4057",
                                font=("TkDefaultFont", 11, "bold"),
                                cursor="question_arrow")
        mintok_info.pack(side="left", padx=(8, 0))
        Tooltip(mintok_info, "Shortest segment (in words) to process. Lower "
                             "values keep short terms from lists and tables; "
                             "higher values focus on running prose and cut "
                             "short fragments. "
                             "Leave blank to use the greyed-out default.\n\n"
                             "Not applied to Chinese and Japanese, which are "
                             "written without spaces: every segment is "
                             "processed, since even a one-character segment "
                             "can hold a term.")

        # -- minimum count (same row as the junk filter) --
        line2 = adv_junk
        ttk.Label(line2, text="Min. count:").pack(side="left", padx=(20, 6))
        self._mincount_is_placeholder = {}
        self.min_count_labels = {}
        for n in NGRAM_CHOICES:
            lbl = ttk.Label(line2, text=f"{n}g")
            lbl.pack(side="left", padx=(4, 1))
            self.min_count_labels[n] = lbl
            e = ttk.Entry(line2, width=3)
            e.pack(side="left")
            e.insert(0, str(self.min_count_defaults[n]))
            e.config(foreground="grey")
            self._mincount_is_placeholder[n] = True
            e.bind("<FocusIn>", lambda ev, n=n: self._on_mincount_focus_in(n))
            e.bind("<FocusOut>", lambda ev, n=n: self._on_mincount_focus_out(n))
            e.bind("<KeyRelease>", lambda ev, n=n: self._on_mincount_key(n))
            self.min_count_entries[n] = e

        info = ttk.Label(line2, text="ⓘ", foreground="#2E4057",
                         font=("TkDefaultFont", 11, "bold"), cursor="question_arrow")
        info.pack(side="left", padx=(10, 0))
        Tooltip(info, "Minimum number of segments a term must appear in to be "
                      "kept, set separately per n-gram length. A segment is "
                      "one TU (TMX) or one sentence (PDF/DOCX).\n\n"
                      "The default is 1, i.e. no frequency filtering, because "
                      "a meaningful threshold depends entirely on how big your "
                      "corpus is. Prefer the 'Max terms' cap for limiting the "
                      "list size; raise these values only if you know your "
                      "corpus well. Leave a field blank to use the greyed-out "
                      "default.\n\n"
                      "For Chinese and Japanese, the first field (labelled "
                      "'all') applies to every term length.")

        # -- stopwords line --
        line3 = ttk.Frame(adv_frame)
        line3.pack(fill="x", padx=8, pady=(2, 6))
        ttk.Label(line3, text="Custom stopwords:").pack(side="left", padx=(0, 6))
        # Pack the info icon on the right first so the entry expands beside it.
        stop_info = ttk.Label(line3, text="ⓘ", foreground="#2E4057",
                              font=("TkDefaultFont", 11, "bold"),
                              cursor="question_arrow")
        stop_info.pack(side="right", padx=(8, 0))
        Tooltip(stop_info, "Words to exclude from results (e.g. a company or "
                           "product name). Separate with commas; matching is "
                           "case-insensitive. A candidate containing any of "
                           "these words is dropped entirely.\n\n"
                           "For Chinese and Japanese, a stopword must match a "
                           "whole word as spaCy segments it, so a stopword "
                           "inside a longer compound may not be caught.")
        self.stopwords_entry = ttk.Entry(line3)
        self.stopwords_entry.pack(side="left", fill="x", expand=True)
        self.STOPWORDS_PLACEHOLDER = "add words to exclude, separated by commas"
        self.stopwords_entry.insert(0, self.STOPWORDS_PLACEHOLDER)
        self.stopwords_entry.config(foreground="grey")
        self.stopwords_entry.bind("<FocusIn>", self._on_stopwords_focus_in)
        self.stopwords_entry.bind("<FocusOut>", self._on_stopwords_focus_out)

        # Row 2b: AI enhancement panel
        ai_frame = ttk.LabelFrame(self.root, text="AI enhancement",
                                  style="Section.TLabelframe")
        ai_frame.pack(fill="x", **pad)

        # -- checkboxes line --
        ai_line0 = ttk.Frame(ai_frame)
        ai_line0.pack(fill="x", padx=8, pady=(6, 2))
        self.ai_clean_check = ttk.Checkbutton(
            ai_line0, text="Clean", variable=self.ai_clean_var,
            command=self._on_ai_clean_toggle)
        self.ai_clean_check.pack(side="left", padx=(0, 12))
        self.ai_translate_check = ttk.Checkbutton(
            ai_line0, text="Translate + define",
            variable=self.ai_translate_var,
            command=self._on_ai_translate_toggle)
        self.ai_translate_check.pack(side="left", padx=(0, 0))
        steps_info = ttk.Label(ai_line0, text="ⓘ", foreground="#2E4057",
                               font=("TkDefaultFont", 11, "bold"),
                               cursor="question_arrow")
        steps_info.pack(side="left", padx=(8, 0))
        Tooltip(steps_info, "'Clean' sends the whole candidate list to the "
                            "AI in one request, which merges variants, "
                            "normalises forms and drops junk. It is the "
                            "master switch for the AI panel.\n\n"
                            "'Translate + define' then processes the cleaned "
                            "terms in batches: each request returns a target "
                            "term and a short definition. Terms with "
                            "human-verified TMX translations have their "
                            "target term read from the TMX rather than "
                            "translated from scratch.")

        # -- model + API key line --
        ai_line_model = ttk.Frame(ai_frame)
        ai_line_model.pack(fill="x", padx=8, pady=2)
        ttk.Label(ai_line_model, text="Model:").pack(side="left", padx=(0, 6))
        self.ai_model_combo = ttk.Combobox(
            ai_line_model, textvariable=self.ai_model_var,
            values=list(AI_MODELS.keys()), state="readonly", width=18)
        self.ai_model_combo.pack(side="left")
        self.ai_model_combo.bind("<<ComboboxSelected>>", self._on_model_change)

        ai_line1 = ai_line_model
        ttk.Label(ai_line1, text="API key:").pack(side="left", padx=(20, 6))
        self.api_key_entry = ttk.Entry(ai_line1, textvariable=self.api_key_var,
                                       show="•")
        self.api_key_entry.pack(side="left", fill="x", expand=True)
        # Any manual keystroke marks the key as user-owned, so switching models
        # won't overwrite it with an environment-variable value.
        self.api_key_entry.bind("<KeyRelease>", self._on_api_key_edit)
        key_info = ttk.Label(ai_line1, text="ⓘ", foreground="#2E4057",
                             font=("TkDefaultFont", 11, "bold"),
                             cursor="question_arrow")
        key_info.pack(side="left", padx=(8, 0))
        Tooltip(key_info, "API key for the selected model's provider "
                          "(Anthropic for Claude models, OpenAI for GPT "
                          "models). Used only for this session and never saved "
                          "by the tool. Pre-filled from ANTHROPIC_API_KEY or "
                          "OPENAI_API_KEY to match the selected model.")

        # -- target + definition language line --
        ai_line2 = ttk.Frame(ai_frame)
        ai_line2.pack(fill="x", padx=8, pady=2)
        ttk.Label(ai_line2, text="Target language:").pack(side="left", padx=(0, 6))
        self.target_language_entry = ttk.Entry(
            ai_line2, textvariable=self.target_language_var, width=20)
        self.target_language_entry.pack(side="left")
        tgt_info = ttk.Label(ai_line2, text="ⓘ", foreground="#2E4057",
                             font=("TkDefaultFont", 11, "bold"),
                             cursor="question_arrow")
        tgt_info.pack(side="left", padx=(8, 0))
        Tooltip(tgt_info, "Language the AI translates the terms into (e.g. "
                          "'Italian'). Fills {{TARGET_LANGUAGE}} in the "
                          "translation prompts.")

        ttk.Label(ai_line2, text="Definition language:").pack(side="left",
                                                              padx=(20, 6))
        self.definition_language_entry = ttk.Entry(
            ai_line2, textvariable=self.definition_language_var, width=20)
        self.definition_language_entry.pack(side="left")
        def_info = ttk.Label(ai_line2, text="ⓘ", foreground="#2E4057",
                             font=("TkDefaultFont", 11, "bold"),
                             cursor="question_arrow")
        def_info.pack(side="left", padx=(8, 0))
        Tooltip(def_info, "Language the AI writes the definitions in (e.g. "
                          "'English' or 'Italian'). Used by the 'Translate + "
                          "define' step; fills {{DEFINITION_LANGUAGE}} in the "
                          "translation prompts.")

        # -- domain line --
        ai_line3 = ttk.Frame(ai_frame)
        ai_line3.pack(fill="x", padx=8, pady=2)
        ttk.Label(ai_line3, text="Domain:").pack(side="left", padx=(0, 6))
        self.domain_entry = ttk.Entry(ai_line3)
        self.domain_entry.pack(side="left", fill="x", expand=True)
        dom_info = ttk.Label(ai_line3, text="ⓘ", foreground="#2E4057",
                             font=("TkDefaultFont", 11, "bold"),
                             cursor="question_arrow")
        dom_info.pack(side="left", padx=(8, 0))
        Tooltip(dom_info, "Short domain label for the glossary, e.g. 'medical "
                          "device' or 'automotive'. Fills {{DOMAIN}} in the "
                          "prompts.")

        # -- context line (multi-line) --
        ai_line4 = ttk.Frame(ai_frame)
        ai_line4.pack(fill="x", padx=8, pady=(2, 6))
        ctx_label_col = ttk.Frame(ai_line4)
        ctx_label_col.pack(side="left", anchor="n")
        ttk.Label(ctx_label_col, text="Context:").pack(side="top", padx=(0, 6))
        ctx_info = ttk.Label(ctx_label_col, text="ⓘ", foreground="#2E4057",
                             font=("TkDefaultFont", 11, "bold"),
                             cursor="question_arrow")
        ctx_info.pack(side="top", padx=(0, 6), pady=(4, 0))
        Tooltip(ctx_info, "A sentence or two describing the source material "
                          "(what it is, its purpose, audience). Fills "
                          "{{CONTEXT}} in the prompts. Free text. Drag the "
                          "grip at the bottom-right to resize.")

        # Container holding the text+scrollbar and, beneath them, a drag handle
        # that resizes only the box (in text lines), pushing the widgets below
        # it (Run button, log) down as it grows.
        ctx_box_col = ttk.Frame(ai_line4)
        ctx_box_col.pack(side="left", fill="both", expand=True, padx=(6, 0))

        text_row = ttk.Frame(ctx_box_col)
        text_row.pack(side="top", fill="both", expand=True)
        default_font = tkfont.nametofont("TkDefaultFont")
        self.context_text = tk.Text(text_row, height=4, wrap="word",
                                    font=default_font, undo=True)
        self.context_text.pack(side="left", fill="both", expand=True)
        ctx_scroll = ttk.Scrollbar(text_row, orient="vertical",
                                   command=self.context_text.yview)
        ctx_scroll.pack(side="right", fill="y")
        self.context_text.config(yscrollcommand=ctx_scroll.set)

        # Drag handle: a thin strip the user can drag vertically to change the
        # box height. Dragging updates the Text's 'height' (in lines), so the
        # box grows/shrinks and everything packed below it shifts accordingly.
        self._ctx_handle = ttk.Separator(ctx_box_col, orient="horizontal")
        self._ctx_handle.pack(side="top", fill="x", pady=(2, 0))
        self._ctx_handle.configure(cursor="sb_v_double_arrow")
        self._ctx_drag_anchor = None
        self._ctx_handle.bind("<Button-1>", self._ctx_resize_start)
        self._ctx_handle.bind("<B1-Motion>", self._ctx_resize_drag)

        # Row 3: run button
        run_frame = ttk.Frame(self.root)
        run_frame.pack(fill="x", **pad)
        self.run_btn = ttk.Button(run_frame, text="Run", command=self._on_run)
        self.run_btn.pack(side="left")

        # Row 4: status line
        self.status = tk.StringVar(value="Browse to a folder, set options, then Run.")
        ttk.Label(self.root, textvariable=self.status).pack(fill="x", padx=10)

        # Row 5: log box
        log_frame = ttk.LabelFrame(self.root, text="Log")
        log_frame.pack(fill="both", expand=True, **pad)
        self.log_box = scrolledtext.ScrolledText(log_frame, state="disabled",
                                                  wrap="word", height=18)
        self.log_box.pack(fill="both", expand=True, padx=6, pady=6)
        # Red styling for error/warning lines (e.g. truncated AI responses).
        self.log_box.tag_config("error", foreground="#B00020")

        # Set the initial enabled/disabled state of the whole AI panel to match
        # the checkboxes (both start on).
        self._refresh_ai_states()

    # ---- limit field placeholder handling ----
    def _on_limit_focus_in(self, event):
        if self.limit_entry.get() == self.LIMIT_PLACEHOLDER:
            self.limit_entry.delete(0, "end")
            self.limit_entry.config(foreground="black")

    def _on_limit_focus_out(self, event):
        if not self.limit_entry.get().strip():
            self.limit_entry.delete(0, "end")
            self.limit_entry.insert(0, self.LIMIT_PLACEHOLDER)
            self.limit_entry.config(foreground="grey")

    def _on_limit_key(self, event):
        current = self.limit_entry.get()
        if current == self.LIMIT_PLACEHOLDER:
            return
        digits = "".join(ch for ch in current if ch.isdigit())
        if digits != current:
            self.limit_entry.delete(0, "end")
            self.limit_entry.insert(0, digits)
            self.limit_entry.config(foreground="black")

    def _get_limit(self) -> int | None:
        val = self.limit_entry.get().strip()
        if val == self.LIMIT_PLACEHOLDER or not val:
            return None
        try:
            n = int(val)
            return n if n > 0 else None
        except ValueError:
            return None

    def _get_selected_ngrams(self) -> set[int]:
        return {n for n, var in self.ngram_vars.items() if var.get()}

    # ---- min-count field placeholder handling ----
    def _on_mincount_focus_in(self, n):
        e = self.min_count_entries[n]
        if self._mincount_is_placeholder[n]:
            e.delete(0, "end")
            e.config(foreground="black")
            self._mincount_is_placeholder[n] = False

    def _on_mincount_focus_out(self, n):
        e = self.min_count_entries[n]
        if not e.get().strip():
            e.delete(0, "end")
            e.insert(0, str(self.min_count_defaults[n]))
            e.config(foreground="grey")
            self._mincount_is_placeholder[n] = True

    def _on_mincount_key(self, n):
        e = self.min_count_entries[n]
        # Focus-in already cleared any placeholder, so here we only keep the
        # field to digits and ensure the text is shown in black.
        current = e.get()
        digits = "".join(ch for ch in current if ch.isdigit())
        if digits != current:
            e.delete(0, "end")
            e.insert(0, digits)
        if current:
            e.config(foreground="black")
            self._mincount_is_placeholder[n] = False

    def _get_min_count(self) -> dict[int, int]:
        """Return the min-count dict, using defaults for blank/placeholder fields."""
        result = {}
        for n, e in self.min_count_entries.items():
            val = e.get().strip()
            if self._mincount_is_placeholder[n] or not val:
                result[n] = self.min_count_defaults[n]
            else:
                try:
                    result[n] = max(1, int(val))
                except ValueError:
                    result[n] = self.min_count_defaults[n]
        return result

    # ---- min-segment-tokens field placeholder handling ----
    def _on_mintok_focus_in(self, event):
        if self._mintok_is_placeholder:
            self.mintok_entry.delete(0, "end")
            self.mintok_entry.config(foreground="black")
            self._mintok_is_placeholder = False

    def _on_mintok_focus_out(self, event):
        if not self.mintok_entry.get().strip():
            self.mintok_entry.delete(0, "end")
            self.mintok_entry.insert(0, str(self.min_segment_tokens_default))
            self.mintok_entry.config(foreground="grey")
            self._mintok_is_placeholder = True

    def _on_mintok_key(self, event):
        e = self.mintok_entry
        current = e.get()
        digits = "".join(ch for ch in current if ch.isdigit())
        if digits != current:
            e.delete(0, "end")
            e.insert(0, digits)
        if current:
            e.config(foreground="black")
            self._mintok_is_placeholder = False

    def _get_min_segment_tokens(self) -> int:
        val = self.mintok_entry.get().strip()
        if self._mintok_is_placeholder or not val:
            return self.min_segment_tokens_default
        try:
            return max(1, int(val))
        except ValueError:
            return self.min_segment_tokens_default

    # ---- stopwords field placeholder handling ----
    def _on_stopwords_focus_in(self, event):
        if self.stopwords_entry.get() == self.STOPWORDS_PLACEHOLDER:
            self.stopwords_entry.delete(0, "end")
            self.stopwords_entry.config(foreground="black")

    def _on_stopwords_focus_out(self, event):
        if not self.stopwords_entry.get().strip():
            self.stopwords_entry.delete(0, "end")
            self.stopwords_entry.insert(0, self.STOPWORDS_PLACEHOLDER)
            self.stopwords_entry.config(foreground="grey")

    def _get_stopwords(self) -> set:
        val = self.stopwords_entry.get().strip()
        if val == self.STOPWORDS_PLACEHOLDER or not val:
            return set()
        return {w.strip().lower() for w in val.split(",") if w.strip()}

    # ---- AI enhancement handling ----
    @staticmethod
    def _env_key_for_label(model_label: str) -> str:
        """Return the API key from the environment variable matching the
        model's provider: OPENAI_API_KEY for GPT models, ANTHROPIC_API_KEY for
        Claude models. Empty string if the variable isn't set."""
        entry = AI_MODELS.get(model_label)
        provider = entry[1] if entry else "anthropic"
        var = "OPENAI_API_KEY" if provider == "openai" else "ANTHROPIC_API_KEY"
        return os.environ.get(var, "")

    def _on_model_change(self, event=None):
        """When the model changes, re-seed the API key from the matching
        provider's environment variable — unless the user has typed their own
        key, in which case leave it alone."""
        if self._api_key_user_edited:
            return
        self.api_key_var.set(self._env_key_for_label(self.ai_model_var.get()))

    def _on_api_key_edit(self, event=None):
        self._api_key_user_edited = True
    def _set_widget_state(self, widget, enabled: bool):
        """Enable/disable a single AI widget, using the right state token for
        its type (comboboxes must return to 'readonly', not 'normal')."""
        if enabled:
            state = "readonly" if isinstance(widget, ttk.Combobox) else "normal"
        else:
            state = "disabled"
        widget.config(state=state)

    def _refresh_ai_states(self):
        """Central state logic for the AI panel.

        'Clean' is the master switch. When off, every dependent control is
        greyed out and 'Translate + define' is unticked. When on, 'Translate +
        define' is enabled, and the target and definition language fields are
        gated on it.
        """
        clean_on = self.ai_clean_var.get()

        # Translate + define depends on cleaning; untick it when cleaning is off.
        self._set_widget_state(self.ai_translate_check, clean_on)
        if not clean_on:
            self.ai_translate_var.set(False)

        # Controls gated purely on cleaning.
        for w in (self.ai_model_combo, self.api_key_entry,
                  self.domain_entry, self.context_text):
            self._set_widget_state(w, clean_on)

        # Both language fields are gated on cleaning AND translate + define.
        translate_on = clean_on and self.ai_translate_var.get()
        self._set_widget_state(self.target_language_entry, translate_on)
        self._set_widget_state(self.definition_language_entry, translate_on)

    def _on_ai_clean_toggle(self):
        # Re-ticking 'Clean' restores 'Translate + define' to its default (on),
        # then applies gating.
        if self.ai_clean_var.get():
            self.ai_translate_var.set(True)
        self._refresh_ai_states()

    def _on_ai_translate_toggle(self):
        self._refresh_ai_states()

    # ---- context box drag-to-resize (in text lines) ----
    def _ctx_resize_start(self, event):
        # Anchor the drag: remember where the pointer grabbed and the box's
        # current height in lines.
        self._ctx_drag_anchor = (event.y_root, int(self.context_text["height"]))

    def _ctx_resize_drag(self, event):
        if self._ctx_drag_anchor is None:
            return
        start_y, start_lines = self._ctx_drag_anchor
        # Convert pixel movement into lines using the font's line height.
        line_px = max(1, tkfont.nametofont("TkDefaultFont").metrics("linespace"))
        delta_lines = (event.y_root - start_y) // line_px
        new_height = int(start_lines + delta_lines)
        new_height = max(2, min(new_height, 40))   # clamp to a sane range
        if new_height != int(self.context_text["height"]):
            self.context_text.config(height=new_height)

    def _get_context(self) -> str:
        return self.context_text.get("1.0", "end").strip()

    def _refresh_language_states(self):
        """Grey out the word-based options for languages written without
        spaces (Chinese, Japanese): n-gram checkboxes, Min. segment length,
        the 'per n-gram' cap, and the 2g-5g Min. count fields (the first
        field becomes the single value for all lengths)."""
        unspaced = self.language_var.get() in UNSPACED_LANGUAGES
        word_based = [self.mintok_entry, self.per_ngram_radio]
        word_based += list(self.ngram_checks.values())
        word_based += [e for n, e in self.min_count_entries.items() if n != 1]
        word_based += [l for n, l in self.min_count_labels.items() if n != 1]
        for w in word_based:
            if isinstance(w, ttk.Label):
                w.config(foreground="grey" if unspaced else "")
            else:
                self._set_widget_state(w, not unspaced)
        self.min_count_labels[1].config(text="all" if unspaced else "1g")
        if unspaced:
            self.limit_mode_var.set("total")

    # ---- actions ----
    def _on_browse(self):
        path = filedialog.askdirectory(title="Select folder to scan")
        if path:
            self.selected_path.set(path)

    def _on_run(self):
        path = self.selected_path.get().strip()
        if not path:
            self._append_error("Please select a folder first.")
            return
        if self.worker and self.worker.is_alive():
            return

        selected_ngrams = self._get_selected_ngrams()
        if (not selected_ngrams
                and self.language_var.get() not in UNSPACED_LANGUAGES):
            self._append_error("Please select at least one n-gram length.")
            return
        term_limit = self._get_limit()
        custom_stopwords = self._get_stopwords()
        min_count = self._get_min_count()
        language = self.language_var.get()
        min_segment_tokens = self._get_min_segment_tokens()

        # --- AI settings + validation ---
        ai_clean = self.ai_clean_var.get()
        ai_translate = self.ai_translate_var.get() and ai_clean
        api_key = self.api_key_var.get().strip()
        _model_entry = AI_MODELS.get(self.ai_model_var.get())
        ai_model = _model_entry[0] if _model_entry else ""
        domain = self.domain_entry.get().strip()
        context = self._get_context()
        target_language = self.target_language_var.get().strip()
        definition_language = self.definition_language_var.get().strip()

        if ai_clean:
            if not api_key:
                self._append_error("AI cleaning is enabled but no API key was "
                                   "provided. Enter a key or turn off AI "
                                   "enhancement.")
                return
            if ai_translate and not target_language:
                self._append_error("Translate + define is enabled but no "
                                   "target language was provided. Enter a "
                                   "target language or turn off Translate + "
                                   "define.")
                return
            if ai_translate and not definition_language:
                self._append_error("Translate + define is enabled but no "
                                   "definition language was provided. Enter "
                                   "a definition language or turn off "
                                   "Translate + define.")
                return

        # One logger for the whole session: buffers everything (precheck +
        # pipeline) so it can be saved to a file at the end. Precheck runs on
        # the main thread and logs straight to the box; the pipeline (worker
        # thread) routes through the queue.
        logger = Logger(sink=self._append_log, error_sink=self._append_error)

        tmx_target_codes = None
        preflight_usage = _new_usage_totals()
        # --- Discovery + TMX srclang pre-check (on main thread, so we can prompt) ---
        by_ext = discover_files(path, logger, self.recurse_var.get())
        tmx_paths = by_ext[".tmx"]
        pdf_paths = by_ext[".pdf"]
        docx_paths = by_ext[".docx"]

        if not (tmx_paths or pdf_paths or docx_paths):
            logger.log("No supported files found. Nothing to do.")
            return

        compatible_tmx, problematic_tmx = check_tmx_source_languages(tmx_paths, language)

        # Log the srclang found in every TMX file (compatible or not)
        if tmx_paths:
            logger.log("[tmx] Source languages found:")
            for p in tmx_paths:
                sl = read_tmx_srclang(p)
                logger.log(f"[tmx]   {p.name}: srclang={sl or 'missing'}")

        if problematic_tmx:
            proceed = self._prompt_problematic_tmx(problematic_tmx, language)
            if not proceed:
                logger.log("Process stopped by user. No files were processed.")
                return
            # User chose to skip problematic TMX files
            logger.log(f"[tmx] Skipping {len(problematic_tmx)} incompatible "
                       f"TMX file(s); proceeding with the rest.")
            tmx_paths = compatible_tmx
            # If skipping leaves nothing at all, stop early
            if not (tmx_paths or pdf_paths or docx_paths):
                logger.log("Nothing left to process after skipping. Stopping.")
                return
        elif tmx_paths:
            logger.log(f"[tmx] All {len(tmx_paths)} TMX file(s) match the "
                       f"selected language ({language}).")

        # --- TMX target-language resolution (pre-flight) ---
        # With the translate step ON, the requested target language is matched
        # against each file's codes via one AI call (and a dialog if needed).
        # With translate OFF there is no requested language, so the user picks
        # directly from the codes found — no AI involved.
        if tmx_paths and ai_translate:
            outcome = self._resolve_tmx_targets(
                tmx_paths, target_language, api_key, ai_model, logger,
                preflight_usage)
            if outcome is None:
                logger.log("Process stopped by user during target-language "
                           "resolution. No files were processed.")
                return
            tmx_target_codes = outcome
        elif tmx_paths:
            logger.log("[tmx] Selecting target languages for the target "
                       "example columns...")
            outcome = self._resolve_tmx_targets_manual(tmx_paths, logger)
            if outcome is None:
                logger.log("Process stopped by user during target-language "
                           "selection. No files were processed.")
                return
            tmx_target_codes = outcome

        self._set_busy(True)
        # From here the pipeline runs in a worker thread, so switch the logger
        # to route through the thread-safe queue while keeping the same buffer.
        logger.sink = lambda m: self.log_queue.put(("log", m))
        logger.update_sink = lambda m: self.log_queue.put(("update", m))
        logger.error_sink = lambda m: self.log_queue.put(("error", m))

        def task():
            try:
                run_pipeline(path, language, tmx_paths, pdf_paths, docx_paths,
                             selected_ngrams, term_limit, custom_stopwords,
                             min_count, min_segment_tokens, logger,
                             ai_clean=ai_clean, ai_translate=ai_translate,
                             api_key=api_key, ai_model=ai_model,
                             domain=domain, context=context,
                             target_language=target_language,
                             definition_language=definition_language,
                             tmx_target_codes=tmx_target_codes,
                             preflight_usage=preflight_usage,
                             limit_mode=self.limit_mode_var.get(),
                             junk_filter=self.junk_filter_var.get())
            except Exception as e:
                logger.log(f"\n!!! Error: {e}")
            finally:
                # Save the full session log next to the source folder.
                try:
                    stamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
                    folder = Path(path)
                    log_name = f"{folder.name}_{stamp}.log"
                    logger.save_to_file(str(folder / log_name))
                    self.log_queue.put(("log", Logger._stamp(
                        f"[log] Session log saved to '{log_name}'.")))
                except Exception as e:
                    self.log_queue.put(("log", Logger._stamp(
                        f"[log] Could not save session log: {e}")))
                self.log_queue.put(("done", None))

        self.worker = threading.Thread(target=task, daemon=True)
        self.worker.start()

    def _prompt_problematic_tmx(self, problematic, language) -> bool:
        """
        Show a modal dialog listing TMX files whose source language doesn't
        match the selected language. Return True to proceed (skipping them),
        False to stop the whole process.
        """
        dialog = tk.Toplevel(self.root)
        dialog.title("Source language mismatch")
        dialog.transient(self.root)
        dialog.grab_set()          # modal
        dialog.geometry("620x420")

        msg = (f"The following TMX file(s) have a source language that does not "
               f"match the selected language ({language}), or an ambiguous/missing "
               f"source language.\n\nYou can skip these files and process the rest, "
               f"or stop to check them separately.")
        ttk.Label(dialog, text=msg, wraplength=580, justify="left").pack(
            padx=12, pady=(12, 6), anchor="w")

        list_frame = ttk.Frame(dialog)
        list_frame.pack(fill="both", expand=True, padx=12, pady=6)
        box = scrolledtext.ScrolledText(list_frame, wrap="none", height=12)
        box.pack(fill="both", expand=True)
        for p, reason in problematic:
            box.insert("end", f"• {p.name}   [{reason}]\n")
        box.config(state="disabled")

        result = {"proceed": False}

        def do_skip():
            result["proceed"] = True
            dialog.destroy()

        def do_stop():
            result["proceed"] = False
            dialog.destroy()

        btns = ttk.Frame(dialog)
        btns.pack(fill="x", padx=12, pady=(6, 12))
        ttk.Button(btns, text="Skip these files and proceed",
                   command=do_skip).pack(side="right", padx=(6, 0))
        ttk.Button(btns, text="Stop",
                   command=do_stop).pack(side="right")

        dialog.protocol("WM_DELETE_WINDOW", do_stop)  # closing window = stop
        self.root.wait_window(dialog)
        return result["proceed"]

    def _prompt_tmx_target_resolution(self, multi_files, no_match_files,
                                      target_language):
        """
        Combined pre-flight dialog for TMX target-language resolution.

        multi_files:    list of (path, [codes]) where several codes matched the
                        requested target language and the user must pick one.
        no_match_files: list of path with no compatible target language.

        Returns None if the user cancels/stops the whole run, otherwise a dict
        {str(path): chosen_code_or_None} covering ONLY the files shown here
        (single-match files are resolved by the caller). A value of None means
        "use as source-only".
        """
        dialog = tk.Toplevel(self.root)
        dialog.title("TMX target language")
        dialog.transient(self.root)
        dialog.grab_set()
        dialog.geometry("640x520")

        intro = (f"Some TMX files need a decision about their target language "
                 f"for '{target_language}'. Review the choices below, then "
                 f"continue or stop.")
        ttk.Label(dialog, text=intro, wraplength=600, justify="left").pack(
            padx=12, pady=(12, 6), anchor="w")

        body = ttk.Frame(dialog)
        body.pack(fill="both", expand=True, padx=12, pady=6)

        SKIP_LABEL = "Skip (use as source-only)"
        combo_by_file = {}

        # Section 1: multi-match files -> a dropdown each.
        if multi_files:
            ttk.Label(body, text="Files with more than one matching target "
                                 "language — choose one per file:",
                      justify="left").pack(anchor="w", pady=(4, 4))
            for p, codes in multi_files:
                row = ttk.Frame(body)
                row.pack(fill="x", pady=2)
                ttk.Label(row, text=p.name, width=32).pack(side="left")
                var = tk.StringVar(value=codes[0])
                combo = ttk.Combobox(row, textvariable=var, state="readonly",
                                     values=list(codes) + [SKIP_LABEL], width=28)
                combo.pack(side="left", padx=(6, 0))
                combo_by_file[str(p)] = var

        # Section 2: no-match files -> listed, resolved by the overall choice.
        if no_match_files:
            ttk.Label(body, text="Files with no target segments in "
                                 f"'{target_language}' (these can only be used "
                                 "as source-only):", justify="left").pack(
                anchor="w", pady=(12, 4))
            nm_box = scrolledtext.ScrolledText(body, wrap="none", height=6)
            nm_box.pack(fill="both", expand=True)
            for p in no_match_files:
                nm_box.insert("end", f"• {p.name}\n")
            nm_box.config(state="disabled")

        result = {"map": None}

        def do_continue():
            resolved = {}
            for key, var in combo_by_file.items():
                val = var.get()
                resolved[key] = None if val == SKIP_LABEL else val
            # No-match files always become source-only when proceeding.
            for p in no_match_files:
                resolved[str(p)] = None
            result["map"] = resolved
            dialog.destroy()

        def do_stop():
            result["map"] = None
            dialog.destroy()

        btns = ttk.Frame(dialog)
        btns.pack(fill="x", padx=12, pady=(6, 12))
        proceed_text = ("Continue" if not no_match_files
                        else "Continue (no-match files as source-only)")
        ttk.Button(btns, text=proceed_text, command=do_continue).pack(
            side="right", padx=(6, 0))
        ttk.Button(btns, text="Stop", command=do_stop).pack(side="right")

        dialog.protocol("WM_DELETE_WINDOW", do_stop)   # closing = stop
        self.root.wait_window(dialog)
        return result["map"]

    def _resolve_tmx_targets(self, tmx_paths, target_language, api_key,
                             ai_model, logger, usage_totals):
        """
        Pre-flight orchestration for TMX target-language resolution. Returns:
          - a dict {str(path): code_or_None} to proceed with, or
          - None if the user stopped the run (or a hard error made it
            impossible to proceed).

        Files with exactly one matching code are auto-resolved. Files with
        several matches, or none, are decided via one combined dialog.
        """
        logger.log("[tmx] Resolving target languages for the translate step...")
        # Stage 1: collect per-file target codes.
        codes_by_file = {}
        for p in tmx_paths:
            codes = collect_tmx_target_codes(p)
            codes_by_file[str(p)] = codes
            logger.log(f"[tmx]   {p.name}: target codes = "
                       f"{', '.join(codes) if codes else '(none)'}")

        # Stage 2: one LLM call to match the user's target against the codes.
        try:
            prompts = load_prompts(logger)
            client = make_ai_client(api_key, _model_provider(ai_model))
            resolution = resolve_tmx_target_codes(
                codes_by_file, target_language,
                prompts["tmx_language_resolution"], client, ai_model, logger,
                usage_totals=usage_totals)
        except Exception as e:
            logger.log_error(
                f"[tmx] Could not resolve TMX target languages "
                f"({type(e).__name__}: {e}). Cannot proceed with translation.")
            return None

        # Sort files into single / multi / no-match.
        auto = {}                 # path -> chosen code (single match)
        multi_files = []          # (Path, [codes])
        no_match_files = []       # Path
        for p in tmx_paths:
            matches = resolution.get(str(p), {}).get("matches", [])
            if len(matches) == 1:
                auto[str(p)] = matches[0]
            elif len(matches) >= 2:
                multi_files.append((p, matches))
            else:
                no_match_files.append(p)

        # If nothing needs a decision, we're done.
        final_map = dict(auto)
        if not multi_files and not no_match_files:
            logger.log("[tmx] All TMX files resolved to a single target "
                       "language automatically.")
            for k, v in final_map.items():
                logger.log(f"[tmx]   {Path(k).name} -> {v}")
            return final_map

        # Otherwise, ask the user (combined dialog).
        decided = self._prompt_tmx_target_resolution(
            multi_files, no_match_files, target_language)
        if decided is None:
            return None   # user stopped
        final_map.update(decided)

        # Log the final decisions, including source-only files.
        for p in tmx_paths:
            code = final_map.get(str(p))
            if code:
                logger.log(f"[tmx]   {p.name} -> target '{code}'")
            else:
                logger.log(f"[tmx]   {p.name} -> source-only (no verified target)")
        if all(final_map.get(str(p)) is None for p in tmx_paths):
            logger.log_error(
                "[tmx] No TMX file will provide verified translations; all are "
                "source-only. Proceeding as chosen.")
        return final_map

    def _prompt_tmx_manual_target(self, multi_files):
        """
        AI-independent dialog: let the user choose one target language code per
        TMX file that declares several. Used when the translation step is OFF
        (so no requested target language exists to match against).

        multi_files: list of (path, [codes]).
        Returns None if the user cancels/stops, otherwise
        {str(path): chosen_code_or_None}.
        """
        dialog = tk.Toplevel(self.root)
        dialog.title("TMX target language")
        dialog.transient(self.root)
        dialog.grab_set()
        dialog.geometry("640x420")

        intro = ("Some TMX files contain more than one target language. Choose "
                 "which one to use for the target examples in each file. "
                 "(Files with a single target language are used automatically.)")
        ttk.Label(dialog, text=intro, wraplength=600, justify="left").pack(
            padx=12, pady=(12, 8), anchor="w")

        body = ttk.Frame(dialog)
        body.pack(fill="both", expand=True, padx=12, pady=4)

        NONE_LABEL = "None — no target examples"
        combo_by_file = {}
        for p, codes in multi_files:
            row = ttk.Frame(body)
            row.pack(fill="x", pady=3)
            ttk.Label(row, text=p.name, width=32).pack(side="left")
            var = tk.StringVar(value=codes[0])
            combo = ttk.Combobox(row, textvariable=var, state="readonly",
                                 values=list(codes) + [NONE_LABEL], width=30)
            combo.pack(side="left", padx=(6, 0))
            combo_by_file[str(p)] = var

        result = {"map": None}

        def do_continue():
            resolved = {}
            for key, var in combo_by_file.items():
                val = var.get()
                resolved[key] = None if val == NONE_LABEL else val
            result["map"] = resolved
            dialog.destroy()

        def do_stop():
            result["map"] = None
            dialog.destroy()

        btns = ttk.Frame(dialog)
        btns.pack(fill="x", padx=12, pady=(8, 12))
        ttk.Button(btns, text="Continue", command=do_continue).pack(
            side="right", padx=(6, 0))
        ttk.Button(btns, text="Stop", command=do_stop).pack(side="right")
        dialog.protocol("WM_DELETE_WINDOW", do_stop)
        self.root.wait_window(dialog)
        return result["map"]

    def _resolve_tmx_targets_manual(self, tmx_paths, logger):
        """
        AI-independent target-language resolution, used when the translation
        step is OFF. Scans each TMX file's target-language codes:
          - exactly one code  -> resolved automatically (no prompt)
          - two or more codes -> the user picks one (or none) in a dialog
          - no codes          -> no targets available for that file
        Returns {str(path): code_or_None}, or None if the user stopped.
        """
        codes_by_file = {}
        for p in tmx_paths:
            codes_by_file[p] = collect_tmx_target_codes(p)

        final_map = {}
        multi_files = []
        for p, codes in codes_by_file.items():
            if len(codes) == 1:
                final_map[str(p)] = codes[0]
            elif len(codes) >= 2:
                multi_files.append((p, codes))
            else:
                final_map[str(p)] = None

        if multi_files:
            decided = self._prompt_tmx_manual_target(multi_files)
            if decided is None:
                return None
            final_map.update(decided)

        for p in tmx_paths:
            code = final_map.get(str(p))
            if code:
                logger.log(f"[tmx]   {p.name} -> target '{code}'")
            else:
                logger.log(f"[tmx]   {p.name} -> no target examples")
        return final_map

    def _set_busy(self, busy: bool):
        if busy:
            self.run_btn.config(state="disabled")
            self.browse_btn.config(state="disabled")
            self.recurse_check.config(state="disabled")
            self.status.set("Processing…")
        else:
            self.run_btn.config(state="normal")
            self.browse_btn.config(state="normal")
            self.recurse_check.config(state="normal")
            self.status.set("Done. Ready for another run.")

    # ---- log box ----
    def _append_log(self, msg: str):
        self.log_box.config(state="normal")
        self.log_box.insert("end", msg + "\n")
        self.log_box.see("end")
        self.log_box.config(state="disabled")
        self._last_line_updatable = False

    def _append_error(self, msg: str):
        """Append a message rendered in red (uses the 'error' text tag)."""
        self.log_box.config(state="normal")
        self.log_box.insert("end", msg + "\n", "error")
        self.log_box.see("end")
        self.log_box.config(state="disabled")
        self._last_line_updatable = False

    def _update_log(self, msg: str):
        self.log_box.config(state="normal")
        if self._last_line_updatable:
            self.log_box.delete("end-2l", "end-1l")
        self.log_box.insert("end", msg + "\n")
        self.log_box.see("end")
        self.log_box.config(state="disabled")
        self._last_line_updatable = True

    def _poll_log_queue(self):
        try:
            while True:
                kind, msg = self.log_queue.get_nowait()
                if kind == "done":
                    self._set_busy(False)
                elif kind == "update":
                    self._update_log(msg)
                elif kind == "error":
                    self._append_error(msg)
                else:
                    self._append_log(msg)
        except queue.Empty:
            pass
        self.root.after(100, self._poll_log_queue)


def main():
    root = tk.Tk()
    App(root)
    root.mainloop()


if __name__ == "__main__":
    main()
